"""Collection: workbooks."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from uuid import uuid4

import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill

from tcc_crypto.collection.config import CollectionConfig
from tcc_crypto.collection.metadata import _git_revision, _selection_source_sha256
from tcc_crypto.collection.schema import DECIMAL_TEXT_COLUMNS
from tcc_crypto.collection.validation import nonstandard_candle_durations


def _format_worksheet(worksheet: openpyxl.worksheet.worksheet.Worksheet) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    worksheet.sheet_view.showGridLines = False


def write_pair_workbook(
    output_path: Path,
    frame: pd.DataFrame,
    metadata: pd.DataFrame,
    validation: pd.DataFrame,
    missing: pd.DatetimeIndex,
    config: CollectionConfig | None = None,
) -> None:
    """Write a derived review workbook without replacing an existing file."""
    if output_path.exists():
        raise FileExistsError(f"Refusing to overwrite workbook: {output_path}")
    fatal_checks = validation.loc[validation["Status"] == "FAIL"]
    if not fatal_checks.empty:
        failed_names = ", ".join(fatal_checks["Check"].astype(str))
        raise ValueError(f"Fatal validation checks failed: {failed_names}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = output_path.with_name(
        f".{output_path.stem}.{uuid4().hex}.tmp.xlsx"
    )
    export_frame = frame.copy()
    for column_name in DECIMAL_TEXT_COLUMNS:
        export_frame[column_name] = pd.to_numeric(
            export_frame[column_name], errors="raise"
        )
    export_frame["open_time_utc"] = export_frame["open_time_utc"].dt.tz_localize(None)
    export_frame["close_time_utc"] = export_frame["close_time_utc"].dt.tz_localize(None)
    missing_frame = pd.DataFrame(
        {"Missing Open Time (UTC)": pd.DatetimeIndex(missing).tz_localize(None)}
    )
    duration_frame = nonstandard_candle_durations(frame, config)
    duration_frame["open_time_utc"] = duration_frame["open_time_utc"].dt.tz_localize(
        None
    )
    duration_frame["close_time_utc"] = duration_frame["close_time_utc"].dt.tz_localize(
        None
    )

    try:
        with pd.ExcelWriter(temporary_path, engine="openpyxl") as writer:
            export_frame.to_excel(writer, sheet_name="Klines", index=False)
            metadata.to_excel(writer, sheet_name="Metadata", index=False)
            validation.to_excel(writer, sheet_name="Validation", index=False)
            missing_frame.to_excel(writer, sheet_name="Missing Intervals", index=False)
            duration_frame.to_excel(
                writer,
                sheet_name="Nonstandard Durations",
                index=False,
            )
            for worksheet in writer.book.worksheets:
                _format_worksheet(worksheet)

            klines_sheet = writer.book["Klines"]
            klines_sheet.column_dimensions["A"].width = 17
            klines_sheet.column_dimensions["B"].width = 22
            for column_letter in [
                "C",
                "D",
                "E",
                "F",
                "G",
                "H",
                "I",
                "J",
                "K",
                "L",
                "M",
                "N",
            ]:
                klines_sheet.column_dimensions[column_letter].width = 20
            for cell in klines_sheet["B"][1:]:
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            close_time_column = export_frame.columns.get_loc("close_time_utc") + 1
            for column in klines_sheet.iter_cols(
                min_col=close_time_column,
                max_col=close_time_column,
                min_row=2,
            ):
                for cell in column:
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"

            writer.book["Metadata"].column_dimensions["A"].width = 36
            writer.book["Metadata"].column_dimensions["B"].width = 100
            writer.book["Validation"].column_dimensions["A"].width = 42
            writer.book["Validation"].column_dimensions["B"].width = 18
            writer.book["Validation"].column_dimensions["C"].width = 18
            writer.book["Validation"].column_dimensions["D"].width = 12
            writer.book["Missing Intervals"].column_dimensions["A"].width = 26
            duration_sheet = writer.book["Nonstandard Durations"]
            duration_sheet.column_dimensions["A"].width = 22
            duration_sheet.column_dimensions["B"].width = 22

        if output_path.exists():
            raise FileExistsError(f"Refusing to overwrite workbook: {output_path}")
        temporary_path.rename(output_path)
    finally:
        temporary_path.unlink(missing_ok=True)


def _validate_workbook(
    path: Path,
    expected_rows: int,
    config: CollectionConfig,
) -> None:
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise ValueError(f"Existing workbook cannot be opened: {path}") from error
    try:
        required_sheets = {
            "Klines",
            "Metadata",
            "Validation",
            "Missing Intervals",
            "Nonstandard Durations",
        }
        if not required_sheets.issubset(workbook.sheetnames):
            raise ValueError(f"Existing workbook is missing required sheets: {path}")
        if workbook["Klines"].max_row - 1 != expected_rows:
            raise ValueError(
                f"Existing workbook row count differs from raw CSV: {path}"
            )
        metadata_values = {
            row[0]: row[1]
            for row in workbook["Metadata"].iter_rows(
                min_row=2,
                max_col=2,
                values_only=True,
            )
            if row[0] is not None
        }
        if (
            metadata_values.get("Requested end (UTC)")
            != config.end_time_utc.isoformat()
        ):
            raise ValueError(f"Existing workbook has a different requested end: {path}")
    finally:
        workbook.close()


def _write_manifest(
    results: pd.DataFrame,
    config: CollectionConfig,
    repository_root: Path,
    execution_time: datetime,
    execution_symbols: tuple[str, ...],
) -> Path:
    output_directory = repository_root / config.output_directory
    output_directory.mkdir(parents=True, exist_ok=True)
    timestamp = execution_time.strftime("%Y%m%dT%H%M%S%fZ")
    output_path = output_directory / f"collection_manifest_{timestamp}.xlsx"
    methodology = pd.DataFrame(
        [
            ("Collection ID", config.collection_id),
            ("Execution timestamp (UTC)", execution_time.isoformat()),
            ("Requested start (UTC)", config.start_time_utc.isoformat()),
            ("Requested end (UTC)", config.end_time_utc.isoformat()),
            ("Interval", config.interval),
            ("Selected pairs", len(config.symbols)),
            ("Execution pair count", len(execution_symbols)),
            ("Execution pairs", ", ".join(execution_symbols)),
            ("Selection source", config.selection_source),
            (
                "Selection source SHA-256",
                _selection_source_sha256(config, repository_root),
            ),
            ("Selection rule", config.selection_rule),
            ("Source endpoint", f"{config.base_url}{config.kline_path}"),
            ("Canonical raw format", "gzip-compressed CSV"),
            ("Code revision", _git_revision(repository_root)),
        ],
        columns=["Field", "Value"],
    )
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        results.to_excel(writer, sheet_name="Collection Results", index=False)
        methodology.to_excel(writer, sheet_name="Methodology", index=False)
        for worksheet in writer.book.worksheets:
            _format_worksheet(worksheet)
            for column in worksheet.columns:
                letter = column[0].column_letter
                max_length = max(len(str(cell.value or "")) for cell in column[:100])
                worksheet.column_dimensions[letter].width = min(
                    max(max_length + 2, 12),
                    80,
                )
    return output_path
