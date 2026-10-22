"""Excel inspection copies of saved analysis tables, with native Excel charts."""

from __future__ import annotations

from math import ceil
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.axis import DateAxis
from openpyxl.chart.data_source import NumFmt

from tcc_crypto.features.catalog import feature_definitions
from tcc_crypto.paths import ANALYSIS_WORKBOOK_NAME
from tcc_crypto.reporting.analysis_definitions import analysis_definitions
from tcc_crypto.reporting.descriptive_definitions import definition_rows
from tcc_crypto.reporting.tables import ensure_writable
from tcc_crypto.reporting.workbooks import configuration_table, write_workbook
from tcc_crypto.workflows.exploration_calculations import ExplorationResult
from tcc_crypto.workflows.exploration_context import ExplorationContext
from tcc_crypto.workflows.exploration_inputs import ExplorationInputs


def _native_daily_charts(path: Path) -> None:
    """Attach native charts using the already written chart-support values."""
    book = load_workbook(path)
    data = book["Chart Data"]
    for row in data.iter_rows(min_row=2, min_col=1, max_col=1):
        cell = row[0]
        if isinstance(cell.value, str):
            cell.value = pd.Timestamp(cell.value).to_pydatetime()
        cell.number_format = "yyyy-mm-dd"
    overview = book["Overview"]
    for title, columns, anchor, unit in [
        ("Daily opening and closing prices", (2, 3), "A12", "USDT per base-asset unit"),
        ("Daily quote volume", (4,), "A29", "USDT"),
    ]:
        chart = LineChart()
        chart.title = title
        chart.y_axis.title = unit
        chart.y_axis.numFmt = NumFmt(formatCode="0.##########", sourceLinked=False)
        chart.x_axis = DateAxis(crossAx=100)
        chart.x_axis.number_format = "yyyy-mm"
        first_date, last_date = data.cell(2, 1).value, data.cell(data.max_row, 1).value
        months = (
            (last_date.year - first_date.year) * 12
            + last_date.month
            - first_date.month
            + 1
        )
        chart.x_axis.majorTimeUnit = "months"
        chart.x_axis.majorUnit = max(1, ceil(months / 8))
        chart.x_axis.title = "UTC date"
        chart.display_blanks = "gap"
        for column in columns:
            chart.add_data(
                Reference(
                    data,
                    min_col=column,
                    max_col=column,
                    min_row=1,
                    max_row=data.max_row,
                ),
                titles_from_data=True,
            )
        chart.set_categories(
            Reference(data, min_col=1, min_row=2, max_row=data.max_row)
        )
        chart.width = 25
        chart.height = 9
        overview.add_chart(chart, anchor)
    book.save(path)
    book.close()


def _select_symbol(table: pd.DataFrame, symbol: str) -> pd.DataFrame:
    column = "Symbol" if "Symbol" in table else "symbol"
    return table.loc[table[column] == symbol].copy() if column in table else table


def _daily_sheets(
    context: ExplorationContext,
    inputs: ExplorationInputs,
    result: ExplorationResult,
    symbol: str,
) -> dict[str, pd.DataFrame]:
    values = inputs.values[symbol]
    chart = values[["open", "close", "quote_asset_volume"]].reindex(
        pd.date_range(values.index.min(), values.index.max(), freq="D")
    )
    chart.index.name = "UTC date"
    chart = chart.reset_index()
    # Excel serial dates are UTC without timezone information; the column explicitly says UTC.
    chart["UTC date"] = (
        chart["UTC date"]
        .map(lambda value: value.tz_localize(None).to_pydatetime())
        .astype(object)
    )
    coverage = _select_symbol(result.tables["coverage"], symbol)
    definitions = pd.DataFrame(definition_rows(context.settings))
    metadata = configuration_table(
        {
            "Symbol": symbol,
            "Analysis": context.paths.analysis_id,
            "Run": context.paths.run_id,
            "Coverage": coverage.to_dict("records"),
        }
    ).rename(columns={"Setting": "Item", "Value": "Definition"})
    return {
        "Overview": _select_symbol(result.tables["statistics"], symbol),
        "Semester Statistics": _select_symbol(result.tables["semesters"], symbol),
        "Extrema": _select_symbol(result.tables["extrema"], symbol),
        "Chart Data": chart,
        "Definitions": pd.concat([metadata, definitions], ignore_index=True),
    }


def _feature_sheets(
    context: ExplorationContext,
    inputs: ExplorationInputs,
    result: ExplorationResult,
    symbol: str,
) -> dict[str, pd.DataFrame]:
    tables = result.tables
    manifest = _select_symbol(tables["manifest"], symbol)
    sheets = {"Summary": manifest}
    key = "features__" + symbol
    if key in tables:
        sheets.update(
            {
                "Features": tables[key].reset_index(),
                "Selection frequency": _select_symbol(
                    tables["selection_frequency"], symbol
                ),
                "Window summary": _select_symbol(tables["window_summary"], symbol),
                "Selection history": _select_symbol(
                    tables["selection_history"], symbol
                ),
            }
        )
    else:
        sheets["Missing dates"] = inputs.missing_dates.get(symbol, pd.DataFrame())
    sheets["Definitions"] = pd.DataFrame(feature_definitions())
    return sheets


def export_workbooks(
    context: ExplorationContext,
    inputs: ExplorationInputs,
    result: ExplorationResult,
    *,
    overwrite: bool = False,
) -> dict[str, Path]:
    """Export only supplied tables; no API requests, feature calculation or model fit."""
    per_symbol = context.analysis in ("daily-descriptive", "feature-selection")
    paths = (
        {
            symbol: context.paths.workbooks / (symbol + ".xlsx")
            for symbol in context.symbols
        }
        if per_symbol
        else {"analysis": context.paths.workbooks / ANALYSIS_WORKBOOK_NAME}
    )
    ensure_writable(list(paths.values()), overwrite=overwrite)
    for key, path in paths.items():
        if context.analysis == "daily-descriptive":
            sheets = _daily_sheets(context, inputs, result, key)
        elif context.analysis == "feature-selection":
            sheets = _feature_sheets(context, inputs, result, key)
        else:
            sheets = standard_sheets(context, inputs, result)
        write_workbook(path, sheets, overwrite=overwrite)
        if context.analysis == "daily-descriptive":
            _native_daily_charts(path)
    return paths


def standard_sheets(
    context: ExplorationContext, inputs: ExplorationInputs, result: ExplorationResult
) -> dict[str, pd.DataFrame]:
    """Keep established worksheet names while adding execution provenance."""
    tables = result.tables
    definitions = analysis_definitions(context, inputs)
    mapping = {
        "coverage": {"coverage": "Kline Coverage"},
        "annual-volume": {
            "annual": "Annual Volume",
            "monthly": "Monthly Data",
            "validation": "Validation",
        },
        "trade-counts": {
            "trade_counts": "Pair Trade Counts",
            "validation": "Validation",
        },
        "hourly-prices": {
            "statistics": "Price Statistics",
            "coverage": "Coverage",
            "extrema": "Extrema",
        },
        "hourly-activity": {
            "statistics": "Activity Statistics",
            "coverage": "Coverage",
            "extrema": "Extrema",
        },
    }
    if context.analysis == "future-high":
        sheets = {"Summary": tables["summary"].reset_index()}
        for symbol in tables["outcomes"]["symbol"].drop_duplicates():
            sheets[symbol] = _select_symbol(tables["outcomes"], symbol)
        sheets["Eligibility"] = tables["eligibility"]
        sheets["Definitions"] = definitions
        return sheets
    titles = mapping.get(
        context.analysis, {name: name.replace("_", " ").title()[:31] for name in tables}
    )
    sheets = {
        title: tables[name].reset_index()
        if not isinstance(tables[name].index, pd.RangeIndex)
        else tables[name]
        for name, title in titles.items()
    }
    sheets[
        "Methodology"
        if context.analysis in ("coverage", "annual-volume", "trade-counts", "earliest")
        else "Definitions"
    ] = definitions
    return sheets
