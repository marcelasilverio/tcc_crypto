"""Human-readable Excel copies of existing pipeline tables; no calculations."""

from __future__ import annotations

import json
import math
from pathlib import Path
from typing import Mapping
from uuid import uuid4

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from tcc_crypto.features.catalog import feature_definitions
from tcc_crypto.reporting.tables import ensure_writable


def _excel_value(value: object) -> object:
    if value is None or value is pd.NA or value is pd.NaT:
        return None
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, default=str)
    if isinstance(value, pd.Timestamp):
        return value.isoformat()
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def write_workbook(
    path: Path, sheets: Mapping[str, pd.DataFrame], *, overwrite: bool = False
) -> Path:
    """Write styled tables with UTC ISO text, filters, and frozen headers."""
    ensure_writable([path], overwrite=overwrite)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, frame in sheets.items():
        if len(frame) + 1 > 1_048_576 or len(frame.columns) > 16_384:
            raise ValueError(
                f"{title} exceeds Excel sheet limits; CSV remains the complete source."
            )
        sheet = workbook.create_sheet(title)
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "B2"
        sheet.append([str(column) for column in frame.columns])
        for row in frame.itertuples(index=False, name=None):
            sheet.append([_excel_value(value) for value in row])
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 32
        for cell in sheet[1]:
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Calibri", size=11)
                cell.alignment = Alignment(vertical="top")
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                elif isinstance(cell.value, float):
                    cell.number_format = "0.0000000000"
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="EDF3F8")
        for column in range(1, len(frame.columns) + 1):
            length = max(
                (
                    len(str(sheet.cell(row, column).value or ""))
                    for row in range(1, min(sheet.max_row, 60) + 1)
                ),
                default=12,
            )
            sheet.column_dimensions[get_column_letter(column)].width = min(
                max(length + 2, 16), 62
            )
        if title in ("Summary", "Definitions", "Configuration"):
            for row in range(2, sheet.max_row + 1):
                sheet.row_dimensions[row].height = 46
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.stem}.{uuid4().hex}.tmp.xlsx")
    try:
        workbook.save(temporary)
        temporary.replace(path)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)
    return path


def configuration_table(configuration: Mapping[str, object]) -> pd.DataFrame:
    return pd.DataFrame(
        [(key, _excel_value(value)) for key, value in configuration.items()],
        columns=["Setting", "Value"],
    )


def write_dataset_workbook(
    path: Path,
    tables: Mapping[str, pd.DataFrame],
    summary: Mapping[str, object],
    *,
    overwrite: bool = False,
) -> Path:
    definitions = pd.DataFrame(feature_definitions())
    target_rows = pd.DataFrame(
        [
            {
                "Feature": f"target_{threshold}pct",
                "Formula": f"1 if high[t+1]/open[t+1]-1 >= {threshold / 100}; 0 otherwise",
                "Inputs": "next UTC day's open and high",
                "Parameters": "missing next day remains unlabeled",
                "Unit": "binary",
                "Implementation": "tcc_crypto.targets.daily.build_next_day_d0_targets",
            }
            for threshold in (2, 3, 4)
        ]
    )
    sheets = {"Summary": configuration_table(summary)}
    for stage, title in (
        ("aligned", "Aligned data"),
        ("features", "Features"),
        ("targets", "Targets"),
        ("dataset", "Modeling dataset"),
    ):
        sheets[title] = tables[stage].reset_index()
    sheets["Definitions"] = pd.concat([definitions, target_rows], ignore_index=True)
    return write_workbook(path, sheets, overwrite=overwrite)


def write_svm_workbook(
    path: Path,
    kind: str,
    tables: Mapping[str, pd.DataFrame],
    configuration: Mapping[str, object],
    *,
    overwrite: bool = False,
) -> Path:
    titles = (
        {"summary": "Metrics", "predictions": "Rolling predictions"}
        if kind == "pilot"
        else {
            "eligibility": "Eligibility",
            "fold_results": "Fold results",
            "candidate_results": "Candidate ranking",
            "selected_configurations": "Selected configurations",
        }
    )
    summary = {
        "Experiment type": kind,
        "Context": configuration["data"],
        "Interpretation": "Classifier evidence only; not portfolio returns or final-test evaluation",
        **{f"{name} rows": len(frame) for name, frame in tables.items()},
    }
    sheets = {
        "Summary": configuration_table(summary),
        **{title: tables[name] for name, title in titles.items()},
        "Configuration": configuration_table(configuration),
    }
    return write_workbook(path, sheets, overwrite=overwrite)
