"""Collect, checkpoint, validate, and export Binance Spot hourly klines.

The canonical raw artifact is a gzip-compressed CSV that preserves the textual
decimal fields returned by Binance. Excel workbooks are derived review copies.
Final raw artifacts are immutable; resumable mutable checkpoints use hidden
files inside each symbol directory until a pair is complete.
"""

from __future__ import annotations

import csv
import gzip
import hashlib
import json
import os
import platform
import shutil
import subprocess
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from datetime import datetime, timezone
from http.client import IncompleteRead, RemoteDisconnected
from pathlib import Path
from typing import Callable, Mapping, Sequence
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen
from uuid import uuid4

import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill

KLINE_COLUMNS = [
    "open_time_ms",
    "open",
    "high",
    "low",
    "close",
    "base_asset_volume",
    "close_time_ms",
    "quote_asset_volume",
    "number_of_trades",
    "taker_buy_base_asset_volume",
    "taker_buy_quote_asset_volume",
    "ignore",
]

DECIMAL_TEXT_COLUMNS = [
    "open",
    "high",
    "low",
    "close",
    "base_asset_volume",
    "quote_asset_volume",
    "taker_buy_base_asset_volume",
    "taker_buy_quote_asset_volume",
]

HOUR_MS = 60 * 60 * 1000


@dataclass(frozen=True)
class CollectionConfig:
    """Validated configuration for one fixed Binance kline collection."""

    collection_id: str
    market: str
    quote_asset: str
    selection_source: str
    selection_snapshot_end_utc: pd.Timestamp
    selection_measure: str
    selection_rule: str
    symbols: tuple[str, ...]
    interval: str
    start_time_utc: pd.Timestamp
    end_time_utc: pd.Timestamp
    base_url: str
    kline_path: str
    time_zone: str
    limit: int
    request_timeout_seconds: int
    max_retries: int
    max_workers: int
    output_directory: Path

    @property
    def start_time_ms(self) -> int:
        """Return the inclusive collection start in Unix milliseconds."""
        return int(self.start_time_utc.timestamp() * 1000)

    @property
    def end_time_ms(self) -> int:
        """Return the inclusive collection end in Unix milliseconds."""
        return int(self.end_time_utc.timestamp() * 1000)


@dataclass(frozen=True)
class PairOutputPaths:
    """Final and resumable artifact paths for one trading pair."""

    directory: Path
    raw_csv_gz: Path
    workbook: Path
    metadata_json: Path
    partial_csv: Path
    checkpoint_json: Path


def _parse_utc_timestamp(value: object, field_name: str) -> pd.Timestamp:
    timestamp = pd.Timestamp(value)
    if timestamp.tzinfo is None:
        raise ValueError(f"{field_name} must include an explicit UTC offset.")
    return timestamp.tz_convert("UTC")


def load_collection_config(path: Path) -> CollectionConfig:
    """Load and validate a JSON collection configuration."""
    with path.open(encoding="utf-8") as config_file:
        payload = json.load(config_file)

    symbols = tuple(str(symbol) for symbol in payload["symbols"])
    if len(symbols) != 40:
        raise ValueError(f"Expected exactly 40 symbols, received {len(symbols)}.")
    if len(set(symbols)) != len(symbols):
        raise ValueError("The configured symbol list contains duplicates.")

    quote_asset = str(payload["quote_asset"])
    if any(not symbol.endswith(quote_asset) for symbol in symbols):
        raise ValueError(f"Every symbol must be quoted in {quote_asset}.")

    interval = str(payload["interval"])
    if interval != "1h":
        raise ValueError("This collector currently validates only the 1h interval.")

    start_time = _parse_utc_timestamp(payload["start_time_utc"], "start_time_utc")
    end_time = _parse_utc_timestamp(payload["end_time_utc"], "end_time_utc")
    if start_time >= end_time:
        raise ValueError("start_time_utc must be earlier than end_time_utc.")
    if start_time.minute or start_time.second or start_time.microsecond:
        raise ValueError("start_time_utc must be aligned to an hourly boundary.")
    expected_end = end_time.normalize() + pd.Timedelta(days=1) - pd.Timedelta(
        milliseconds=1
    )
    if end_time != expected_end:
        raise ValueError("end_time_utc must be the final millisecond of a UTC day.")

    limit = int(payload["limit"])
    if not 1 <= limit <= 1000:
        raise ValueError("Binance kline limit must be between 1 and 1000.")

    return CollectionConfig(
        collection_id=str(payload["collection_id"]),
        market=str(payload["market"]),
        quote_asset=quote_asset,
        selection_source=str(payload["selection_source"]),
        selection_snapshot_end_utc=_parse_utc_timestamp(
            payload["selection_snapshot_end_utc"],
            "selection_snapshot_end_utc",
        ),
        selection_measure=str(payload["selection_measure"]),
        selection_rule=str(payload["selection_rule"]),
        symbols=symbols,
        interval=interval,
        start_time_utc=start_time,
        end_time_utc=end_time,
        base_url=str(payload["base_url"]).rstrip("/"),
        kline_path=str(payload["kline_path"]),
        time_zone=str(payload["time_zone"]),
        limit=limit,
        request_timeout_seconds=int(payload["request_timeout_seconds"]),
        max_retries=int(payload["max_retries"]),
        max_workers=int(payload["max_workers"]),
        output_directory=Path(payload["output_directory"]),
    )


def request_json(
    url: str,
    parameters: Mapping[str, object],
    *,
    timeout_seconds: int,
    max_retries: int,
) -> object:
    """Request public Binance JSON with bounded transient-error retries."""
    request_url = f"{url}?{urlencode(parameters)}"
    for attempt in range(max_retries + 1):
        request = Request(
            request_url,
            headers={"User-Agent": "tcc-crypto-research/0.1"},
        )
        try:
            with urlopen(request, timeout=timeout_seconds) as response:
                return json.load(response)
        except HTTPError as error:
            is_transient = error.code == 429 or 500 <= error.code < 600
            if is_transient and attempt < max_retries:
                retry_after = error.headers.get("Retry-After")
                delay_seconds = (
                    float(retry_after)
                    if retry_after is not None
                    else min(2**attempt, 30)
                )
                time.sleep(delay_seconds)
                continue
            response_body = error.read().decode("utf-8", errors="replace")
            raise RuntimeError(
                f"Binance returned HTTP {error.code} for {url}: {response_body}"
            ) from error
        except URLError as error:
            if attempt < max_retries:
                time.sleep(min(2**attempt, 30))
                continue
            raise RuntimeError(
                f"Could not reach Binance at {url}: {error.reason}"
            ) from error
        except (IncompleteRead, RemoteDisconnected) as error:
            if attempt < max_retries:
                time.sleep(min(2**attempt, 30))
                continue
            raise RuntimeError(
                f"Binance connection was interrupted for {url}: {error}"
            ) from error

    raise RuntimeError(f"Request retry loop ended unexpectedly for {url}.")


PageRequester = Callable[[Mapping[str, object]], object]


def _default_page_requester(config: CollectionConfig) -> PageRequester:
    endpoint = f"{config.base_url}{config.kline_path}"

    def requester(parameters: Mapping[str, object]) -> object:
        return request_json(
            endpoint,
            parameters,
            timeout_seconds=config.request_timeout_seconds,
            max_retries=config.max_retries,
        )

    return requester


def _validate_page(
    payload: object,
    symbol: str,
    next_start_ms: int,
) -> list[list[object]]:
    if not isinstance(payload, list):
        raise ValueError(f"Unexpected kline response type for {symbol}.")
    if not payload:
        return []
    if any(not isinstance(row, list) or len(row) != 12 for row in payload):
        raise ValueError(
            f"Unexpected kline row schema for {symbol}; expected 12 fields."
        )

    rows = [list(row) for row in payload]
    open_times = [int(row[0]) for row in rows]
    if open_times[0] < next_start_ms:
        raise ValueError(f"Kline pagination moved backwards for {symbol}.")
    if any(
        current <= previous
        for previous, current in zip(open_times, open_times[1:])
    ):
        raise ValueError(f"Kline page is not strictly chronological for {symbol}.")
    return rows


def _raw_frame_from_rows(rows: list[list[object]]) -> pd.DataFrame:
    """Create a frame while preserving Binance decimal fields as text."""
    frame = pd.DataFrame(rows, columns=KLINE_COLUMNS)
    frame["open_time_ms"] = pd.to_numeric(
        frame["open_time_ms"], errors="raise", downcast="integer"
    )
    frame["close_time_ms"] = pd.to_numeric(
        frame["close_time_ms"], errors="raise", downcast="integer"
    )
    frame["number_of_trades"] = pd.to_numeric(
        frame["number_of_trades"], errors="raise", downcast="integer"
    )
    for column_name in DECIMAL_TEXT_COLUMNS:
        frame[column_name] = frame[column_name].map(str)
    frame["ignore"] = frame["ignore"].map(str)
    return _add_utc_columns(frame)


def _add_utc_columns(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    if "open_time_utc" not in result.columns:
        result.insert(
            1,
            "open_time_utc",
            pd.to_datetime(result["open_time_ms"], unit="ms", utc=True),
        )
    if "close_time_utc" not in result.columns:
        position = result.columns.get_loc("close_time_ms") + 1
        result.insert(
            position,
            "close_time_utc",
            pd.to_datetime(result["close_time_ms"], unit="ms", utc=True),
        )
    return result


def collect_hourly_klines(
    symbol: str,
    config: CollectionConfig,
    *,
    page_requester: PageRequester | None = None,
) -> pd.DataFrame:
    """Collect one symbol in memory; primarily useful for tests and small runs."""
    requester = page_requester or _default_page_requester(config)
    rows: list[list[object]] = []
    next_start_ms = config.start_time_ms

    while next_start_ms <= config.end_time_ms:
        parameters: dict[str, object] = {
            "symbol": symbol,
            "interval": config.interval,
            "startTime": next_start_ms,
            "endTime": config.end_time_ms,
            "timeZone": config.time_zone,
            "limit": config.limit,
        }
        page = _validate_page(requester(parameters), symbol, next_start_ms)
        if not page:
            break
        rows.extend(page)
        next_start_ms = int(page[-1][0]) + 1
        if len(page) < config.limit:
            break

    if not rows:
        raise ValueError(f"No hourly klines were returned for {symbol}.")
    frame = _raw_frame_from_rows(rows)
    return (
        frame.loc[
            (frame["open_time_ms"] >= config.start_time_ms)
            & (frame["close_time_ms"] <= config.end_time_ms)
        ]
        .sort_values("open_time_ms", kind="stable")
        .reset_index(drop=True)
    )


def pair_output_paths(
    symbol: str,
    config: CollectionConfig,
    repository_root: Path,
) -> PairOutputPaths:
    """Return the structured per-symbol paths for final and partial artifacts."""
    start_date = config.start_time_utc.strftime("%Y-%m-%d")
    end_date = config.end_time_utc.strftime("%Y-%m-%d")
    stem = f"{symbol}_{config.interval}_{start_date}_{end_date}"
    directory = repository_root / config.output_directory / symbol
    return PairOutputPaths(
        directory=directory,
        raw_csv_gz=directory / f"{stem}.csv.gz",
        workbook=directory / f"{stem}.xlsx",
        metadata_json=directory / f"{stem}.metadata.json",
        partial_csv=directory / f".{stem}.partial.csv",
        checkpoint_json=directory / f".{stem}.checkpoint.json",
    )


def pair_output_path(
    symbol: str,
    config: CollectionConfig,
    repository_root: Path,
) -> Path:
    """Return the derived Excel workbook path for compatibility with callers."""
    return pair_output_paths(symbol, config, repository_root).workbook


def _checkpoint_fingerprint(symbol: str, config: CollectionConfig) -> str:
    payload = {
        "symbol": symbol,
        "interval": config.interval,
        "start_time_ms": config.start_time_ms,
        "end_time_ms": config.end_time_ms,
        "base_url": config.base_url,
        "kline_path": config.kline_path,
        "time_zone": config.time_zone,
        "columns": KLINE_COLUMNS,
    }
    serialized = json.dumps(payload, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(serialized.encode("utf-8")).hexdigest()


def _atomic_write_json(path: Path, payload: Mapping[str, object]) -> None:
    temporary_path = path.with_name(f".{path.name}.{uuid4().hex}.tmp")
    try:
        with temporary_path.open("w", encoding="utf-8") as output_file:
            json.dump(payload, output_file, indent=2, sort_keys=True)
            output_file.write("\n")
            output_file.flush()
            os.fsync(output_file.fileno())
        temporary_path.replace(path)
    finally:
        temporary_path.unlink(missing_ok=True)


def _initialize_or_validate_checkpoint(
    symbol: str,
    config: CollectionConfig,
    paths: PairOutputPaths,
) -> tuple[int, int]:
    """Return the next request timestamp and existing checkpoint row count."""
    paths.directory.mkdir(parents=True, exist_ok=True)
    fingerprint = _checkpoint_fingerprint(symbol, config)

    if paths.checkpoint_json.exists():
        with paths.checkpoint_json.open(encoding="utf-8") as checkpoint_file:
            checkpoint = json.load(checkpoint_file)
        if checkpoint.get("fingerprint") != fingerprint:
            raise ValueError(
                f"Checkpoint configuration differs for {symbol}: "
                f"{paths.checkpoint_json}"
            )
    else:
        if paths.partial_csv.exists():
            raise ValueError(
                f"Partial CSV exists without its checkpoint: {paths.partial_csv}"
            )
        _atomic_write_json(
            paths.checkpoint_json,
            {
                "symbol": symbol,
                "fingerprint": fingerprint,
                "rows_written": 0,
                "last_open_time_ms": None,
                "next_start_time_ms": config.start_time_ms,
                "updated_at_utc": datetime.now(timezone.utc).isoformat(),
            },
        )

    if not paths.partial_csv.exists():
        with paths.partial_csv.open("x", encoding="utf-8", newline="") as output_file:
            csv.writer(output_file).writerow(KLINE_COLUMNS)
            output_file.flush()
            os.fsync(output_file.fileno())

    row_count = 0
    last_open_time_ms: int | None = None
    with paths.partial_csv.open(encoding="utf-8", newline="") as input_file:
        reader = csv.reader(input_file)
        header = next(reader, None)
        if header != KLINE_COLUMNS:
            raise ValueError(f"Unexpected checkpoint CSV header for {symbol}.")
        for row in reader:
            if len(row) != len(KLINE_COLUMNS):
                raise ValueError(f"Malformed checkpoint CSV row for {symbol}.")
            open_time_ms = int(row[0])
            if last_open_time_ms is not None and open_time_ms <= last_open_time_ms:
                raise ValueError(
                    f"Checkpoint is not strictly chronological for {symbol}."
                )
            last_open_time_ms = open_time_ms
            row_count += 1

    next_start_ms = (
        config.start_time_ms if last_open_time_ms is None else last_open_time_ms + 1
    )
    return next_start_ms, row_count


def _append_checkpoint_page(
    paths: PairOutputPaths,
    symbol: str,
    config: CollectionConfig,
    rows: list[list[object]],
    previous_row_count: int,
) -> int:
    filtered_rows = [
        row
        for row in rows
        if int(row[0]) >= config.start_time_ms
        and int(row[6]) <= config.end_time_ms
    ]
    if filtered_rows:
        with paths.partial_csv.open("a", encoding="utf-8", newline="") as output_file:
            csv.writer(output_file).writerows(filtered_rows)
            output_file.flush()
            os.fsync(output_file.fileno())
    row_count = previous_row_count + len(filtered_rows)
    last_open_time_ms = int(rows[-1][0])
    _atomic_write_json(
        paths.checkpoint_json,
        {
            "symbol": symbol,
            "fingerprint": _checkpoint_fingerprint(symbol, config),
            "rows_written": row_count,
            "last_open_time_ms": last_open_time_ms,
            "next_start_time_ms": last_open_time_ms + 1,
            "updated_at_utc": datetime.now(timezone.utc).isoformat(),
        },
    )
    return row_count


def collect_to_checkpoint(
    symbol: str,
    config: CollectionConfig,
    paths: PairOutputPaths,
    *,
    page_requester: PageRequester | None = None,
) -> Path:
    """Collect pages into a resumable uncompressed CSV checkpoint."""
    requester = page_requester or _default_page_requester(config)
    next_start_ms, row_count = _initialize_or_validate_checkpoint(
        symbol,
        config,
        paths,
    )

    while next_start_ms <= config.end_time_ms:
        parameters: dict[str, object] = {
            "symbol": symbol,
            "interval": config.interval,
            "startTime": next_start_ms,
            "endTime": config.end_time_ms,
            "timeZone": config.time_zone,
            "limit": config.limit,
        }
        page = _validate_page(requester(parameters), symbol, next_start_ms)
        if not page:
            break
        row_count = _append_checkpoint_page(
            paths,
            symbol,
            config,
            page,
            row_count,
        )
        next_start_ms = int(page[-1][0]) + 1
        if len(page) < config.limit:
            break

    if row_count == 0:
        raise ValueError(f"No hourly klines were returned for {symbol}.")
    return paths.partial_csv


def _load_raw_csv(path: Path, *, compressed: bool) -> pd.DataFrame:
    compression = "gzip" if compressed else None
    frame = pd.read_csv(
        path,
        dtype=str,
        keep_default_na=False,
        compression=compression,
    )
    if list(frame.columns) != KLINE_COLUMNS:
        raise ValueError(f"Unexpected raw CSV schema: {path}")
    frame["open_time_ms"] = pd.to_numeric(frame["open_time_ms"], errors="raise")
    frame["close_time_ms"] = pd.to_numeric(frame["close_time_ms"], errors="raise")
    frame["number_of_trades"] = pd.to_numeric(
        frame["number_of_trades"], errors="raise"
    )
    return _add_utc_columns(frame)


def load_raw_csv_gz(path: Path) -> pd.DataFrame:
    """Load a canonical raw CSV while preserving decimal text columns."""
    return _load_raw_csv(path, compressed=True)


def missing_hourly_intervals(frame: pd.DataFrame) -> pd.DatetimeIndex:
    """Return missing hourly openings inside the observed span."""
    if frame.empty:
        return pd.DatetimeIndex([], tz="UTC")
    observed = pd.DatetimeIndex(frame["open_time_utc"])
    expected = pd.date_range(observed.min(), observed.max(), freq="1h")
    return expected.difference(observed)


def nonstandard_candle_durations(frame: pd.DataFrame) -> pd.DataFrame:
    """Return shortened candles whose close remains inside the nominal hour."""
    open_times = pd.to_numeric(frame["open_time_ms"], errors="raise")
    close_times = pd.to_numeric(frame["close_time_ms"], errors="raise")
    duration_ms = close_times - open_times + 1
    shortened = (duration_ms > 0) & (duration_ms < HOUR_MS)
    diagnostics = frame.loc[
        shortened,
        ["open_time_utc", "close_time_utc", "open_time_ms", "close_time_ms"],
    ].copy()
    diagnostics["duration_ms"] = duration_ms.loc[shortened].astype("int64")
    diagnostics["nominal_duration_ms"] = HOUR_MS
    diagnostics["shortfall_ms"] = HOUR_MS - diagnostics["duration_ms"]
    return diagnostics.reset_index(drop=True)


def validate_hourly_klines(
    frame: pd.DataFrame,
    config: CollectionConfig,
) -> pd.DataFrame:
    """Validate temporal boundaries, schema values, and OHLC consistency."""
    missing = missing_hourly_intervals(frame)
    open_times = pd.to_numeric(frame["open_time_ms"], errors="raise")
    close_times = pd.to_numeric(frame["close_time_ms"], errors="raise")
    number_of_trades = pd.to_numeric(frame["number_of_trades"], errors="raise")
    decimals = frame[DECIMAL_TEXT_COLUMNS].apply(pd.to_numeric, errors="raise")
    time_differences = open_times.diff().dropna()

    high = decimals["high"]
    low = decimals["low"]
    open_price = decimals["open"]
    close_price = decimals["close"]
    invalid_ohlc = (
        (high < low)
        | (high < open_price)
        | (high < close_price)
        | (low > open_price)
        | (low > close_price)
    )
    volume_columns = [
        "base_asset_volume",
        "quote_asset_volume",
        "taker_buy_base_asset_volume",
        "taker_buy_quote_asset_volume",
    ]
    negative_volume = (decimals[volume_columns] < 0).any(axis=1)
    non_positive_ohlc = (decimals[["open", "high", "low", "close"]] <= 0).any(
        axis=1
    )
    duration_ms = close_times - open_times + 1
    close_outside_nominal_window = (duration_ms <= 0) | (duration_ms > HOUR_MS)
    shortened_durations = (duration_ms > 0) & (duration_ms < HOUR_MS)

    checks = [
        ("Rows returned", len(frame), ">= 1", "OK" if len(frame) else "FAIL"),
        (
            "Duplicate open timestamps",
            int(open_times.duplicated().sum()),
            0,
            "OK" if not open_times.duplicated().any() else "FAIL",
        ),
        (
            "Non-increasing timestamp transitions",
            int((time_differences <= 0).sum()),
            0,
            "OK" if not (time_differences <= 0).any() else "FAIL",
        ),
        (
            "Open timestamps off the UTC hour",
            int((open_times % HOUR_MS != 0).sum()),
            0,
            "OK" if not (open_times % HOUR_MS != 0).any() else "FAIL",
        ),
        (
            "Close times outside the nominal candle window",
            int(close_outside_nominal_window.sum()),
            0,
            "OK" if not close_outside_nominal_window.any() else "FAIL",
        ),
        (
            "Shortened candle durations",
            int(shortened_durations.sum()),
            0,
            "OK" if not shortened_durations.any() else "REVIEW",
        ),
        (
            "Rows before requested start",
            int((open_times < config.start_time_ms).sum()),
            0,
            "OK" if not (open_times < config.start_time_ms).any() else "FAIL",
        ),
        (
            "Rows after requested end",
            int((close_times > config.end_time_ms).sum()),
            0,
            "OK" if not (close_times > config.end_time_ms).any() else "FAIL",
        ),
        (
            "Non-positive OHLC prices",
            int(non_positive_ohlc.sum()),
            0,
            "OK" if not non_positive_ohlc.any() else "FAIL",
        ),
        (
            "OHLC consistency violations",
            int(invalid_ohlc.sum()),
            0,
            "OK" if not invalid_ohlc.any() else "FAIL",
        ),
        (
            "Rows with negative volume",
            int(negative_volume.sum()),
            0,
            "OK" if not negative_volume.any() else "FAIL",
        ),
        (
            "Rows with negative trade count",
            int((number_of_trades < 0).sum()),
            0,
            "OK" if not (number_of_trades < 0).any() else "FAIL",
        ),
        (
            "Missing hours within observed span",
            len(missing),
            0,
            "OK" if not len(missing) else "REVIEW",
        ),
    ]
    return pd.DataFrame(checks, columns=["Check", "Actual", "Expected", "Status"])


def _assert_collection_complete(
    symbol: str,
    frame: pd.DataFrame,
    validation: pd.DataFrame,
    config: CollectionConfig,
) -> None:
    fatal_checks = validation.loc[validation["Status"] == "FAIL"]
    if not fatal_checks.empty:
        failed_names = ", ".join(fatal_checks["Check"].astype(str))
        raise ValueError(f"Fatal validation checks failed for {symbol}: {failed_names}")
    last_close_ms = int(pd.to_numeric(frame["close_time_ms"]).max())
    if last_close_ms != config.end_time_ms:
        last_close = pd.to_datetime(last_close_ms, unit="ms", utc=True).isoformat()
        raise ValueError(
            f"{symbol} does not reach the requested end; last close is {last_close}."
        )


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as input_file:
        for chunk in iter(lambda: input_file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _git_revision(repository_root: Path) -> str:
    try:
        result = subprocess.run(
            ["git", "rev-parse", "HEAD"],
            cwd=repository_root,
            check=True,
            capture_output=True,
            text=True,
        )
        return result.stdout.strip()
    except (OSError, subprocess.CalledProcessError):
        return "unavailable"


def _selection_source_sha256(
    config: CollectionConfig,
    repository_root: Path,
) -> str:
    source_path = repository_root / config.selection_source
    return _sha256_file(source_path) if source_path.exists() else "unavailable"


def build_metadata(
    symbol: str,
    frame: pd.DataFrame,
    config: CollectionConfig,
    repository_root: Path,
    collected_at_utc: datetime,
    raw_csv_sha256: str,
) -> pd.DataFrame:
    """Build auditable collection metadata for a pair workbook."""
    missing = missing_hourly_intervals(frame)
    shortened = nonstandard_candle_durations(frame)
    first_open = frame["open_time_utc"].min()
    last_close = frame["close_time_utc"].max()
    records = [
        ("Collection ID", config.collection_id),
        ("Collection timestamp (UTC)", collected_at_utc.isoformat()),
        ("Code revision", _git_revision(repository_root)),
        ("Python version", platform.python_version()),
        ("pandas version", pd.__version__),
        ("openpyxl version", openpyxl.__version__),
        ("Market", config.market),
        ("Trading pair", symbol),
        ("Quote asset", config.quote_asset),
        ("Interval", config.interval),
        ("Timezone", "UTC"),
        ("Requested start (UTC)", config.start_time_utc.isoformat()),
        ("Requested end (UTC)", config.end_time_utc.isoformat()),
        ("First observed open (UTC)", first_open.isoformat()),
        ("Last observed close (UTC)", last_close.isoformat()),
        ("Observed hourly klines", len(frame)),
        ("Missing hours within observed span", len(missing)),
        ("Shortened candle durations", len(shortened)),
        ("Selection source", config.selection_source),
        (
            "Selection source SHA-256",
            _selection_source_sha256(config, repository_root),
        ),
        ("Selection snapshot end (UTC)", config.selection_snapshot_end_utc.isoformat()),
        ("Selection measure", config.selection_measure),
        ("Selection rule", config.selection_rule),
        ("Source endpoint", f"{config.base_url}{config.kline_path}"),
        (
            "Canonical raw format",
            "gzip-compressed CSV with Binance decimal text preserved",
        ),
        ("Canonical raw CSV SHA-256", raw_csv_sha256),
        (
            "Excel role",
            "Derived human-review copy; numeric cells use Excel floating point",
        ),
    ]
    return pd.DataFrame(records, columns=["Field", "Value"])


def _format_worksheet(worksheet: openpyxl.worksheet.worksheet.Worksheet) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    worksheet.sheet_view.showGridLines = False


def write_pair_workbook(
    output_path: Path,
    frame: pd.DataFrame,
    metadata: pd.DataFrame,
    validation: pd.DataFrame,
    missing: pd.DatetimeIndex,
) -> None:
    """Write a derived review workbook without replacing an existing file."""
    if output_path.exists():
        raise FileExistsError(f"Refusing to overwrite workbook: {output_path}")
    fatal_checks = validation.loc[validation["Status"] == "FAIL"]
    if not fatal_checks.empty:
        failed_names = ", ".join(fatal_checks["Check"].astype(str))
        raise ValueError(f"Fatal validation checks failed: {failed_names}")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = output_path.with_name(
        f".{output_path.stem}.{uuid4().hex}.tmp.xlsx"
    )
    export_frame = frame.copy()
    for column_name in DECIMAL_TEXT_COLUMNS:
        export_frame[column_name] = pd.to_numeric(
            export_frame[column_name], errors="raise"
        )
    export_frame["open_time_utc"] = export_frame["open_time_utc"].dt.tz_localize(None)
    export_frame["close_time_utc"] = export_frame["close_time_utc"].dt.tz_localize(None)
    missing_frame = pd.DataFrame(
        {"Missing Open Time (UTC)": pd.DatetimeIndex(missing).tz_localize(None)}
    )
    duration_frame = nonstandard_candle_durations(frame)
    duration_frame["open_time_utc"] = duration_frame[
        "open_time_utc"
    ].dt.tz_localize(None)
    duration_frame["close_time_utc"] = duration_frame[
        "close_time_utc"
    ].dt.tz_localize(None)

    try:
        with pd.ExcelWriter(temporary_path, engine="openpyxl") as writer:
            export_frame.to_excel(writer, sheet_name="Klines", index=False)
            metadata.to_excel(writer, sheet_name="Metadata", index=False)
            validation.to_excel(writer, sheet_name="Validation", index=False)
            missing_frame.to_excel(writer, sheet_name="Missing Intervals", index=False)
            duration_frame.to_excel(
                writer,
                sheet_name="Nonstandard Durations",
                index=False,
            )
            for worksheet in writer.book.worksheets:
                _format_worksheet(worksheet)

            klines_sheet = writer.book["Klines"]
            klines_sheet.column_dimensions["A"].width = 17
            klines_sheet.column_dimensions["B"].width = 22
            for column_letter in [
                "C",
                "D",
                "E",
                "F",
                "G",
                "H",
                "I",
                "J",
                "K",
                "L",
                "M",
                "N",
            ]:
                klines_sheet.column_dimensions[column_letter].width = 20
            for cell in klines_sheet["B"][1:]:
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            close_time_column = export_frame.columns.get_loc("close_time_utc") + 1
            for column in klines_sheet.iter_cols(
                min_col=close_time_column,
                max_col=close_time_column,
                min_row=2,
            ):
                for cell in column:
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"

            writer.book["Metadata"].column_dimensions["A"].width = 36
            writer.book["Metadata"].column_dimensions["B"].width = 100
            writer.book["Validation"].column_dimensions["A"].width = 42
            writer.book["Validation"].column_dimensions["B"].width = 18
            writer.book["Validation"].column_dimensions["C"].width = 18
            writer.book["Validation"].column_dimensions["D"].width = 12
            writer.book["Missing Intervals"].column_dimensions["A"].width = 26
            duration_sheet = writer.book["Nonstandard Durations"]
            duration_sheet.column_dimensions["A"].width = 22
            duration_sheet.column_dimensions["B"].width = 22

        if output_path.exists():
            raise FileExistsError(f"Refusing to overwrite workbook: {output_path}")
        temporary_path.rename(output_path)
    finally:
        temporary_path.unlink(missing_ok=True)


def _compress_partial_csv(paths: PairOutputPaths) -> None:
    if paths.raw_csv_gz.exists():
        raise FileExistsError(f"Refusing to overwrite raw CSV: {paths.raw_csv_gz}")
    temporary_path = paths.raw_csv_gz.with_name(
        f".{paths.raw_csv_gz.name}.{uuid4().hex}.tmp"
    )
    try:
        with paths.partial_csv.open("rb") as source_file:
            with temporary_path.open("xb") as compressed_file:
                with gzip.GzipFile(
                    filename="",
                    mode="wb",
                    fileobj=compressed_file,
                    mtime=0,
                ) as gzip_file:
                    shutil.copyfileobj(source_file, gzip_file)
                compressed_file.flush()
                os.fsync(compressed_file.fileno())
        if paths.raw_csv_gz.exists():
            raise FileExistsError(f"Refusing to overwrite raw CSV: {paths.raw_csv_gz}")
        temporary_path.rename(paths.raw_csv_gz)
    finally:
        temporary_path.unlink(missing_ok=True)


def _validate_workbook(
    path: Path,
    expected_rows: int,
    config: CollectionConfig,
) -> None:
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as error:
        raise ValueError(f"Existing workbook cannot be opened: {path}") from error
    try:
        required_sheets = {
            "Klines",
            "Metadata",
            "Validation",
            "Missing Intervals",
            "Nonstandard Durations",
        }
        if not required_sheets.issubset(workbook.sheetnames):
            raise ValueError(f"Existing workbook is missing required sheets: {path}")
        if workbook["Klines"].max_row - 1 != expected_rows:
            raise ValueError(f"Existing workbook row count differs from raw CSV: {path}")
        metadata_values = {
            row[0]: row[1]
            for row in workbook["Metadata"].iter_rows(
                min_row=2,
                max_col=2,
                values_only=True,
            )
            if row[0] is not None
        }
        if metadata_values.get("Requested end (UTC)") != config.end_time_utc.isoformat():
            raise ValueError(f"Existing workbook has a different requested end: {path}")
    finally:
        workbook.close()


def _write_pair_metadata_json(
    paths: PairOutputPaths,
    metadata: pd.DataFrame,
    validation: pd.DataFrame,
) -> None:
    raw_hash = metadata.set_index("Field").loc[
        "Canonical raw CSV SHA-256",
        "Value",
    ]
    if paths.metadata_json.exists():
        with paths.metadata_json.open(encoding="utf-8") as metadata_file:
            existing = json.load(metadata_file)
        if existing.get("raw_csv_sha256") != raw_hash:
            raise ValueError(
                f"Existing metadata checksum differs: {paths.metadata_json}"
            )
        if existing.get("workbook_sha256") != _sha256_file(paths.workbook):
            raise ValueError(
                f"Existing workbook checksum differs: {paths.metadata_json}"
            )
        recorded_metadata = existing.get("metadata", {})
        selection_hash = metadata.set_index("Field").loc[
            "Selection source SHA-256",
            "Value",
        ]
        if recorded_metadata.get("Selection source SHA-256") != selection_hash:
            raise ValueError(
                f"Selection source checksum differs: {paths.metadata_json}"
            )
        return

    metadata_map = {
        str(row["Field"]): row["Value"] for _, row in metadata.iterrows()
    }
    payload = {
        "metadata": metadata_map,
        "validation": validation.to_dict(orient="records"),
        "raw_csv_sha256": raw_hash,
        "workbook_sha256": _sha256_file(paths.workbook),
    }
    temporary_path = paths.metadata_json.with_name(
        f".{paths.metadata_json.name}.{uuid4().hex}.tmp"
    )
    try:
        with temporary_path.open("x", encoding="utf-8") as output_file:
            json.dump(payload, output_file, indent=2, sort_keys=True, default=str)
            output_file.write("\n")
            output_file.flush()
            os.fsync(output_file.fileno())
        if paths.metadata_json.exists():
            raise FileExistsError(
                f"Refusing to overwrite metadata: {paths.metadata_json}"
            )
        temporary_path.rename(paths.metadata_json)
    finally:
        temporary_path.unlink(missing_ok=True)


def _pair_result(
    symbol: str,
    selection_rank: int,
    status: str,
    frame: pd.DataFrame,
    paths: PairOutputPaths,
) -> dict[str, object]:
    missing = missing_hourly_intervals(frame)
    shortened = nonstandard_candle_durations(frame)
    return {
        "Selection Rank": selection_rank,
        "Trading Pair": symbol,
        "Status": status,
        "Rows": len(frame),
        "Missing Hours": len(missing),
        "Shortened Candles": len(shortened),
        "First Open (UTC)": frame["open_time_utc"].min().isoformat(),
        "Last Close (UTC)": frame["close_time_utc"].max().isoformat(),
        "Raw CSV": str(paths.raw_csv_gz),
        "Raw CSV SHA-256": _sha256_file(paths.raw_csv_gz),
        "Workbook": str(paths.workbook),
        "Workbook SHA-256": _sha256_file(paths.workbook),
        "Metadata": str(paths.metadata_json),
        "Error": None,
    }


def collect_symbol_outputs(
    symbol: str,
    selection_rank: int,
    config: CollectionConfig,
    repository_root: Path,
    *,
    page_requester: PageRequester | None = None,
) -> dict[str, object]:
    """Validate existing outputs or collect and create all pair artifacts."""
    paths = pair_output_paths(symbol, config, repository_root)
    paths.directory.mkdir(parents=True, exist_ok=True)
    collected_at = datetime.now(timezone.utc)

    if paths.raw_csv_gz.exists():
        frame = load_raw_csv_gz(paths.raw_csv_gz)
        validation = validate_hourly_klines(frame, config)
        _assert_collection_complete(symbol, frame, validation, config)
        raw_hash = _sha256_file(paths.raw_csv_gz)
        metadata = build_metadata(
            symbol,
            frame,
            config,
            repository_root,
            collected_at,
            raw_hash,
        )
        missing = missing_hourly_intervals(frame)
        if paths.workbook.exists():
            _validate_workbook(paths.workbook, len(frame), config)
        else:
            write_pair_workbook(paths.workbook, frame, metadata, validation, missing)
        _write_pair_metadata_json(paths, metadata, validation)
        paths.partial_csv.unlink(missing_ok=True)
        paths.checkpoint_json.unlink(missing_ok=True)
        return _pair_result(
            symbol,
            selection_rank,
            "VALIDATED_EXISTING",
            frame,
            paths,
        )

    if paths.workbook.exists():
        raise ValueError(
            "A derived workbook exists without its canonical raw CSV; refusing "
            f"to infer raw data from Excel: {paths.workbook}"
        )

    collect_to_checkpoint(
        symbol,
        config,
        paths,
        page_requester=page_requester,
    )
    frame = _load_raw_csv(paths.partial_csv, compressed=False)
    validation = validate_hourly_klines(frame, config)
    _assert_collection_complete(symbol, frame, validation, config)
    _compress_partial_csv(paths)
    raw_hash = _sha256_file(paths.raw_csv_gz)
    metadata = build_metadata(
        symbol,
        frame,
        config,
        repository_root,
        collected_at,
        raw_hash,
    )
    missing = missing_hourly_intervals(frame)
    write_pair_workbook(paths.workbook, frame, metadata, validation, missing)
    _write_pair_metadata_json(paths, metadata, validation)
    paths.partial_csv.unlink(missing_ok=True)
    paths.checkpoint_json.unlink(missing_ok=True)
    return _pair_result(symbol, selection_rank, "COLLECTED", frame, paths)


def collect_symbol_to_workbook(
    symbol: str,
    selection_rank: int,
    config: CollectionConfig,
    repository_root: Path,
) -> dict[str, object]:
    """Compatibility alias for complete per-symbol artifact collection."""
    return collect_symbol_outputs(symbol, selection_rank, config, repository_root)


def _write_manifest(
    results: pd.DataFrame,
    config: CollectionConfig,
    repository_root: Path,
    execution_time: datetime,
    execution_symbols: tuple[str, ...],
) -> Path:
    output_directory = repository_root / config.output_directory
    output_directory.mkdir(parents=True, exist_ok=True)
    timestamp = execution_time.strftime("%Y%m%dT%H%M%S%fZ")
    output_path = output_directory / f"collection_manifest_{timestamp}.xlsx"
    methodology = pd.DataFrame(
        [
            ("Collection ID", config.collection_id),
            ("Execution timestamp (UTC)", execution_time.isoformat()),
            ("Requested start (UTC)", config.start_time_utc.isoformat()),
            ("Requested end (UTC)", config.end_time_utc.isoformat()),
            ("Interval", config.interval),
            ("Selected pairs", len(config.symbols)),
            ("Execution pair count", len(execution_symbols)),
            ("Execution pairs", ", ".join(execution_symbols)),
            ("Selection source", config.selection_source),
            (
                "Selection source SHA-256",
                _selection_source_sha256(config, repository_root),
            ),
            ("Selection rule", config.selection_rule),
            ("Source endpoint", f"{config.base_url}{config.kline_path}"),
            ("Canonical raw format", "gzip-compressed CSV"),
            ("Code revision", _git_revision(repository_root)),
        ],
        columns=["Field", "Value"],
    )
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        results.to_excel(writer, sheet_name="Collection Results", index=False)
        methodology.to_excel(writer, sheet_name="Methodology", index=False)
        for worksheet in writer.book.worksheets:
            _format_worksheet(worksheet)
            for column in worksheet.columns:
                letter = column[0].column_letter
                max_length = max(len(str(cell.value or "")) for cell in column[:100])
                worksheet.column_dimensions[letter].width = min(
                    max(max_length + 2, 12),
                    80,
                )
    return output_path


def run_collection(
    config: CollectionConfig,
    repository_root: Path,
    *,
    symbols: Sequence[str] | None = None,
    progress_callback: Callable[[str], None] | None = None,
) -> tuple[pd.DataFrame, Path]:
    """Run the full collection or an approved subset and write a manifest."""
    execution_symbols = (
        config.symbols if symbols is None else tuple(str(symbol) for symbol in symbols)
    )
    if not execution_symbols:
        raise ValueError("At least one collection symbol is required.")
    if len(set(execution_symbols)) != len(execution_symbols):
        raise ValueError("The execution symbol list contains duplicates.")
    unknown_symbols = [
        symbol for symbol in execution_symbols if symbol not in config.symbols
    ]
    if unknown_symbols:
        raise ValueError(
            "Execution symbols are not in the approved configuration: "
            + ", ".join(unknown_symbols)
        )

    selection_ranks = {
        symbol: rank for rank, symbol in enumerate(config.symbols, start=1)
    }
    execution_time = datetime.now(timezone.utc)
    results: list[dict[str, object]] = []
    workers = max(1, min(config.max_workers, len(execution_symbols)))

    with ThreadPoolExecutor(max_workers=workers) as executor:
        futures = {
            executor.submit(
                collect_symbol_outputs,
                symbol,
                selection_ranks[symbol],
                config,
                repository_root,
            ): (selection_ranks[symbol], symbol)
            for symbol in execution_symbols
        }
        for completed, future in enumerate(as_completed(futures), start=1):
            rank, symbol = futures[future]
            try:
                result = future.result()
            except Exception as error:
                result = {
                    "Selection Rank": rank,
                    "Trading Pair": symbol,
                    "Status": "FAILED",
                    "Rows": None,
                    "Missing Hours": None,
                    "Shortened Candles": None,
                    "First Open (UTC)": None,
                    "Last Close (UTC)": None,
                    "Raw CSV": None,
                    "Raw CSV SHA-256": None,
                    "Workbook": None,
                    "Workbook SHA-256": None,
                    "Metadata": None,
                    "Error": str(error),
                }
            results.append(result)
            if progress_callback is not None:
                progress_callback(
                    f"Completed {completed}/{len(execution_symbols)}: "
                    f"{symbol} ({result['Status']})"
                )

    results_frame = pd.DataFrame(results).sort_values("Selection Rank").reset_index(
        drop=True
    )
    manifest_path = _write_manifest(
        results_frame,
        config,
        repository_root,
        execution_time,
        execution_symbols,
    )
    return results_frame, manifest_path


def find_repository_root(start: Path | None = None) -> Path:
    """Find the repository root by looking for AGENTS.md and pyproject.toml."""
    current = (start or Path.cwd()).resolve()
    for candidate in (current, *current.parents):
        if (candidate / "AGENTS.md").exists() and (
            candidate / "pyproject.toml"
        ).exists():
            return candidate
    raise FileNotFoundError("Could not locate the repository root.")
