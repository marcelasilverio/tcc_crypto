"""Final temporal boundaries, frozen choices, matched baselines and omissions."""

from dataclasses import replace
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from test_daily_svm_grid_search import _dataset, _selection_config, _small_config

from tcc_crypto.modeling.final_config import (
    FinalEvaluationConfig,
    load_final_evaluation_config,
)
from tcc_crypto.modeling.final_evaluation import (
    MODELS,
    freeze_configurations,
    run_final_evaluation,
    summarize_final_predictions,
)
from tcc_crypto.modeling.ranking import _parameter_values
from tcc_crypto.reporting.final_evaluation import write_final_workbook
from tcc_crypto.reporting.tables import read_table, write_table
from tcc_crypto.visualization.final_evaluation import final_evaluation_figures


@pytest.fixture
def example():
    data = _dataset(24)
    data["target_date_utc"] = data.index + pd.Timedelta(days=1)
    grid = replace(_small_config(), development_end_utc=data.index[10])
    config = FinalEvaluationConfig(
        "test", Path("unused.json"), data.index[12], data.index[-1]
    )
    selected = pd.DataFrame(
        [
            {
                "symbol": "BTCUSDT",
                "target": "target_2pct",
                "window_size": 4,
                "candidate_order": 0,
                "selection_rank": 1,
                "failed_fold_count": 0,
                "successful_fold_count": 2,
                **_parameter_values(
                    {"kernel": "linear", "C": 0.1, "class_weight": None}
                ),
            }
        ]
    )
    eligibility = pd.DataFrame(
        [
            {
                "symbol": "BTCUSDT",
                "target": "target_2pct",
                "window_size": 4,
                "status": "eligible",
            }
        ]
    )
    return data, grid, config, selected, eligibility


def evaluate(example, **kwargs):
    data, grid, config, selected, eligibility = example
    return run_final_evaluation(
        data,
        symbol="BTCUSDT",
        selected=selected,
        eligibility=eligibility,
        grid_config=grid,
        config=config,
        selection_config=_selection_config(),
        feature_columns=("signal", "secondary"),
        **kwargs,
    )


def test_final_windows_keep_gap_and_fit_only_past_rows(example):
    data, _, config, _, _ = example
    result = evaluate(example)
    predictions = result.predictions
    assert len(predictions) == 12 * len(MODELS)
    assert predictions.status.eq("ok").all()
    assert predictions.training_rows.eq(4).all()
    assert predictions.training_end_utc.eq(
        predictions.feature_date_utc - pd.Timedelta(days=2)
    ).all()
    assert predictions.training_target_end_utc.lt(predictions.feature_date_utc).all()
    first = predictions.loc[
        predictions.feature_date_utc.eq(config.feature_start_utc)
    ].iloc[0]
    features = first.selected_features.split("|")
    expected = (
        data.loc[first.training_start_utc : first.training_end_utc, features]
        .mean()
        .to_numpy()
    )
    import json

    np.testing.assert_allclose(json.loads(first.scaler_mean), expected)
    assert set(result.metrics.model) == set(MODELS)
    assert result.metrics.evaluated_dates.nunique() == 1
    assert result.metrics.evaluated_dates.eq(12).all()
    assert not result.roc.empty and not result.precision_recall.empty
    assert result.metrics.loc[
        result.metrics.model.eq("always_negative"), "roc_auc"
    ].iloc[0] == pytest.approx(0.5)


def test_future_features_and_labels_cannot_change_earlier_predictions(example):
    original = evaluate(example)
    data, *rest = example
    altered = data.copy()
    cutoff = data.index[-4]
    altered.loc[cutoff:, "signal"] *= 10000
    altered.loc[cutoff:, "target_2pct"] = 0
    changed = evaluate((altered, *rest))
    columns = [
        "model",
        "feature_date_utc",
        "prediction",
        "decision_score",
        "selected_features",
        "scaler_mean",
        "scaler_scale",
    ]
    pd.testing.assert_frame_equal(
        original.predictions.loc[
            original.predictions.feature_date_utc.lt(cutoff), columns
        ].reset_index(drop=True),
        changed.predictions.loc[
            changed.predictions.feature_date_utc.lt(cutoff), columns
        ].reset_index(drop=True),
    )


def test_single_class_windows_keep_dated_omissions_and_common_samples(example):
    data, *rest = example
    data = data.copy()
    data.loc[: data.index[15], "target_2pct"] = 0
    result = evaluate((data, *rest))
    skipped = result.predictions.loc[
        result.predictions.status.eq("single_class_training")
    ]
    assert not skipped.empty
    assert skipped.prediction.isna().all() and skipped.decision_score.isna().all()
    assert not skipped.comparison_eligible.any()
    assert result.metrics.omitted_dates.gt(0).all()
    assert result.metrics.evaluated_dates.nunique() == 1


def test_frozen_parameters_reject_tampering_and_missing_development_choices(example):
    _, grid, _, selected, eligibility = example
    bad = selected.copy()
    bad.loc[0, "C"] = 999
    with pytest.raises(ValueError, match="differs"):
        freeze_configurations(bad, eligibility, "BTCUSDT", grid)
    with pytest.raises(ValueError, match="no selected"):
        freeze_configurations(selected.iloc[:0], eligibility, "BTCUSDT", grid)
    with pytest.raises(ValueError, match="eligibility"):
        freeze_configurations(selected, eligibility, "LUNAUSDT", grid)


def test_excluded_assets_have_no_fits_or_predictions(example, monkeypatch):
    import tcc_crypto.modeling.final_evaluation as module

    data, grid, config, selected, eligibility = example
    eligibility = eligibility.assign(symbol="LUNAUSDT", status="ineligible")
    monkeypatch.setattr(
        module, "prepare_window", lambda *args: pytest.fail("Unexpected fit")
    )
    result = module.run_final_evaluation(
        data,
        symbol="LUNAUSDT",
        selected=selected.iloc[:0],
        eligibility=eligibility,
        grid_config=grid,
        config=config,
        selection_config=_selection_config(),
        feature_columns=("signal", "secondary"),
    )
    assert result.predictions.empty and result.metrics.empty
    assert result.eligibility.final_test_eligibility_status.eq(
        "no_eligible_development_configuration"
    ).all()


def test_single_class_evaluation_has_no_invented_curve(example):
    result = evaluate(example)
    predictions = result.predictions.copy()
    predictions["actual"] = 0
    metrics, roc, pr = summarize_final_predictions(predictions)
    assert metrics.roc_auc.isna().all()
    assert metrics.metric_diagnostic.eq("single_class_evaluation").all()
    assert roc.empty and pr.empty


def test_final_tables_roundtrip_and_export_without_fits(example, tmp_path, monkeypatch):
    import matplotlib.pyplot as plt
    from openpyxl import load_workbook

    import tcc_crypto.modeling.final_evaluation as module

    result = evaluate(example)
    monkeypatch.setattr(
        module, "prepare_window", lambda *args: pytest.fail("Unexpected fit")
    )
    tables = {}
    for name in (
        "eligibility",
        "predictions",
        "metrics",
        "roc",
        "precision_recall",
        "selected_configurations",
    ):
        path = tmp_path / f"{name}.csv.gz"
        write_table(
            getattr(result, name),
            path,
            root=tmp_path,
            context={"test": True},
            sources=[],
        )
        tables[name] = read_table(path, context={"test": True}, stage="final-test")
        pd.testing.assert_frame_equal(
            tables[name], getattr(result, name), check_freq=False
        )
    workbook = tmp_path / "final.xlsx"
    write_final_workbook(
        workbook,
        tables,
        {
            "evaluation": {"evaluation_id": "test"},
            "full_period": {"data": {"symbol": "BTCUSDT"}},
            "nested": {f"setting_{i}": i for i in range(40)},
        },
    )
    book = load_workbook(workbook, read_only=True)
    assert book["Dated predictions"].max_row == len(result.predictions) + 1
    assert book["Metrics"].max_row == len(MODELS) + 1
    assert "nested.setting_39" in [row[0] for row in book["Configuration"].values]
    assert "Configuration" not in [row[0] for row in book["Summary"].values]
    book.close()
    _, _, config, _, _ = example
    figures = final_evaluation_figures(
        result.roc,
        result.precision_recall.assign(window_size=60),
        symbol="BTCUSDT",
        start=config.feature_start_utc,
        end=config.feature_end_utc,
    )
    assert set(figures) == {"roc", "precision_recall"}
    assert figures["precision_recall"].axes[0].lines
    assert all(
        line.get_drawstyle() == "steps-post"
        for line in figures["precision_recall"].axes[0].lines
    )
    for figure in figures.values():
        plt.close(figure)


def test_load_final_configuration_matches_approved_main_protocol():
    root = Path(__file__).resolve().parents[1]
    config = load_final_evaluation_config(root)
    assert config.feature_start_utc == pd.Timestamp("2025-01-01", tz="UTC")
    assert config.feature_end_utc == pd.Timestamp("2026-08-31", tz="UTC")
    assert config.gap == 1


def test_workflow_persists_and_reexports_without_retraining(tmp_path, monkeypatch):
    import json

    from test_daily_pipeline import context as original_fixture

    from tcc_crypto.paths import ArtifactPaths, FinalArtifactPaths
    from tcc_crypto.workflows import daily_data, svm
    from tcc_crypto.workflows import final_evaluation as workflow

    base = original_fixture.__wrapped__(tmp_path)
    grid = replace(
        _small_config(),
        candidate_symbols=("BTCUSDT",),
        raw_directory=base.paths.raw.parent,
        development_start_utc=base.data.sample_start_utc,
        development_end_utc=pd.Timestamp("2019-07-20", tz="UTC"),
    )
    development = replace(
        base,
        profile="development",
        model=grid,
        data=replace(
            base.data, context_id="development", sample_end_utc=grid.development_end_utc
        ),
        paths=ArtifactPaths(tmp_path, "development", "BTCUSDT", base.paths.raw.parent),
    )
    full = replace(
        base,
        profile="experiment",
        data=replace(base.data, context_id="full"),
        paths=ArtifactPaths(tmp_path, "full", "BTCUSDT", base.paths.raw.parent),
    )
    daily_data.prepare(development)
    daily_data.prepare(full)
    svm.run(development)
    config_path = Path("final.json")
    (tmp_path / config_path).write_text("{}")
    config = FinalEvaluationConfig(
        "final",
        tmp_path / "unused",
        pd.Timestamp("2019-07-22", tz="UTC"),
        full.data.sample_end_utc,
    )
    paths = FinalArtifactPaths(tmp_path, "final", "BTCUSDT")
    monkeypatch.setattr(
        workflow, "_inputs", lambda *args: (config, paths, full, {"test": True})
    )
    outputs = workflow.run(development, config_path=config_path)
    assert outputs["workbook"].is_file()
    manifest = paths.directory / "run_metadata.json"
    metadata = json.loads(manifest.read_text())
    assert metadata["test_used_for_hyperparameter_selection"] is False
    assert metadata["development_data_equality"] == "exact_full_period_subset"
    monkeypatch.setattr(
        workflow,
        "run_final_evaluation",
        lambda *args, **kwargs: pytest.fail("Unexpected training"),
    )
    workflow.export(development, config_path=config_path, overwrite=True)
    workflow.export_figures(development, config_path=config_path, overwrite=True)
    overview = workflow.summarize_universe([development], config_path=config_path)
    assert overview["workbook"].is_file()
    combined = pd.read_csv(overview["metrics"])
    assert set(combined.model) == set(MODELS)
    assert combined.evaluated_dates.nunique() == 1
    with pytest.raises(FileExistsError):
        workflow.run(development, config_path=config_path)


def test_future_outcome_columns_cannot_be_requested_as_features(example):
    data, grid, config, selected, eligibility = example
    with pytest.raises(ValueError, match="Future outcomes"):
        run_final_evaluation(
            data,
            symbol="BTCUSDT",
            selected=selected,
            eligibility=eligibility,
            grid_config=grid,
            config=config,
            selection_config=_selection_config(),
            feature_columns=("signal", "target_2pct"),
        )
