"""Characterization and integration tests for independently executable stages."""

import json
import subprocess
import sys
from dataclasses import replace
from io import StringIO
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from tcc_crypto.features.catalog import ACTIVE_FEATURES, feature_definitions
from tcc_crypto.features.daily import calculate_daily_features
from tcc_crypto.paths import ArtifactPaths
from tcc_crypto.reporting.provenance import file_sha256
from tcc_crypto.reporting.tables import read_table, write_table
from tcc_crypto.targets.daily import build_next_day_d0_targets
from tcc_crypto.workflows import daily_data, svm
from tcc_crypto.workflows.context import load_context

ROOT = Path(__file__).resolve().parents[1]
REFERENCES = Path(__file__).parent / "fixtures" / "daily_refactor"


def test_full_experiment_context_preserves_modeling_boundaries():
    full = load_context("experiment", root=ROOT)
    development = load_context("development", root=ROOT)
    settings = json.loads(
        (ROOT / "configs/binance_daily_svm_grid_search.json").read_text()
    )
    assert full.model == development.model
    assert full.data.sample_start_utc == development.data.sample_start_utc
    assert full.data.warmup_start_utc == pd.Timestamp("2019-06-01", tz="UTC")
    assert full.data.sample_end_utc == pd.Timestamp(
        settings["final_test"]["feature_end_utc"]
    )
    assert full.data.sample_end_utc + pd.Timedelta(days=1) == pd.Timestamp(
        settings["final_test"]["target_end_utc"]
    )
    assert full.paths.context_id != development.paths.context_id
    assert "results_workbook" not in full.paths.inventory(full.kind)
    assert "LUNAUSDT" in full.model.candidate_symbols
    assert "LUNAUSDT" in full.model.excluded_symbols


@pytest.mark.parametrize(
    "stage,extra",
    [
        ("collect", []),
        ("pilot", []),
        ("grid-search", []),
        ("final-test", []),
        ("export", ["--kind", "results"]),
        ("export", ["--kind", "final-results"]),
    ],
)
def test_full_experiment_cli_rejects_collection_and_models(stage, extra, monkeypatch):
    from tcc_crypto.daily import main
    from tcc_crypto.workflows import context as contexts

    monkeypatch.setattr(
        contexts,
        "load_context",
        lambda *args, **kwargs: pytest.fail("Unexpected execution"),
    )
    with pytest.raises(SystemExit) as error:
        main([stage, "--profile", "experiment", *extra])
    assert error.value.code == 2


def test_full_experiment_api_rejects_modeling_before_reading_data():
    full = load_context("experiment", root=ROOT)
    with pytest.raises(ValueError, match="preparation-only"):
        svm.run(full)
    with pytest.raises(ValueError, match="preparation-only"):
        svm.export_results(full)


def assert_reference(name, actual):
    expected = pd.read_json(
        StringIO((REFERENCES / f"{name}.json").read_text()), orient="table"
    )
    actual = actual.drop(
        columns=["fit_seconds", "mean_fit_seconds"], errors="ignore"
    ).copy()
    actual.attrs = {}
    # pandas table JSON infers all-null object columns as float; recover their
    # original types from the pre-refactor schema recorded with the fixtures.
    schemas = json.loads((REFERENCES / "schema.json").read_text())
    from tcc_crypto.reporting.tables import _restore_objects

    for column, dtype in schemas[name]["dtypes"].items():
        if dtype == "object":
            expected[column] = expected[column].astype(object)
    _restore_objects(expected, schemas[name]["objects"])
    pd.testing.assert_frame_equal(
        actual, expected, check_freq=False, check_exact=False, rtol=1e-12, atol=1e-12
    )


def test_feature_and_target_characterization():
    aligned = pd.read_json(
        StringIO((REFERENCES / "aligned.json").read_text()), orient="table"
    )
    features = calculate_daily_features(aligned)
    assert_reference("features", features)
    assert_reference("targets", build_next_day_d0_targets(aligned, features.index))
    assert tuple(row["Feature"] for row in feature_definitions()) == ACTIVE_FEATURES


def test_pilot_and_grid_characterization():
    from test_daily_svm_grid_search import _dataset, _small_config
    from test_daily_svm_pilot import _pilot_dataset, _selection_config

    from tcc_crypto.modeling.grid_search import run_temporal_grid_search
    from tcc_crypto.modeling.pilot import run_default_svc_rolling_pilot

    pilot = run_default_svc_rolling_pilot(
        _pilot_dataset(),
        selection_config=_selection_config(),
        feature_columns=("signal", "secondary"),
        target_columns=("target_2pct",),
    )
    assert_reference("pilot_predictions", pilot.predictions)
    assert_reference("pilot_summary", pilot.summary)
    grid = run_temporal_grid_search(
        _dataset(),
        symbol="BTCUSDT",
        config=_small_config(),
        selection_config=_selection_config(),
        feature_columns=("signal", "secondary"),
    )
    for name in (
        "eligibility",
        "fold_results",
        "candidate_results",
        "selected_configurations",
    ):
        assert_reference("grid_" + name, getattr(grid, name))


@pytest.fixture
def context(tmp_path):
    original = load_context("pilot", root=ROOT)
    raw = tmp_path / "raw" / "BTCUSDT"
    raw.mkdir(parents=True)
    dates = pd.date_range("2019-06-01", periods=85, tz="UTC", freq="D")
    values = pd.DataFrame(
        {
            "open_time_ms": dates.asi8 // 1_000_000,
            "open": [100 + i / 10 for i in range(len(dates))],
            "high": [
                (100 + i / 10) * (1 + (0.01, 0.025, 0.035, 0.055)[i % 4])
                for i in range(len(dates))
            ],
            "low": [99 + i / 10 for i in range(len(dates))],
            "close": [
                100 + i / 10 + (0.1 if i % 2 else -0.1) for i in range(len(dates))
            ],
            "quote_asset_volume": [1000.0 + (i % 7) * 10 for i in range(len(dates))],
        }
    ).drop(index=[35, 36])
    values.to_csv(raw / "BTCUSDT_1d_sample.csv.gz", index=False)
    data = replace(original.data, sample_end_utc=pd.Timestamp("2019-08-20", tz="UTC"))
    selection = replace(original.features.feature_selection, training_window_sizes=(4,))
    model = replace(
        original.model,
        raw_directory=raw.parent,
        sample_end_utc=data.sample_end_utc,
        training_window_sizes=(4,),
    )
    return replace(
        original,
        root=tmp_path,
        data=data,
        model=model,
        features=replace(original.features, feature_selection=selection),
        paths=ArtifactPaths(tmp_path, data.context_id, data.symbol, raw.parent),
    )


def test_stages_equal_in_memory_builder_and_preserve_raw(context):
    from tcc_crypto.workflows.datasets import build_daily_modeling_dataset

    raw = daily_data.raw_source(context)
    checksum = file_sha256(raw)
    expected = build_daily_modeling_dataset(context.model)
    daily_data.prepare(context)
    actual = daily_data.read_stage(context, "dataset")
    actual.attrs = {}
    pd.testing.assert_frame_equal(actual, expected, check_freq=False, check_exact=True)
    assert file_sha256(raw) == checksum
    aligned = daily_data.read_stage(context, "aligned")
    assert aligned.is_imputed.sum() == 2
    assert str(aligned.index.tz) == "UTC"
    features = daily_data.read_stage(context, "features")
    assert features.iloc[0].isna().any()
    assert features.index.max() == context.data.sample_end_utc
    assert actual.target_2pct.dtype == "Int8"
    assert actual.target_is_imputed.dtype == "boolean"
    assert actual.feature_is_imputed.dtype == "boolean"
    with pytest.raises(FileExistsError, match="overwrite"):
        daily_data.prepare(context)
    daily_data.prepare(context, overwrite=True)


def test_missing_stale_and_modified_inputs_are_rejected(context):
    with pytest.raises(FileNotFoundError, match="align"):
        daily_data.features(context)
    daily_data.prepare(context)
    modified = replace(
        context,
        data=replace(
            context.data,
            sample_end_utc=context.data.sample_end_utc - pd.Timedelta(days=1),
        ),
    )
    with pytest.raises(ValueError, match="Incompatible"):
        daily_data.read_stage(modified, "features")
    path = context.paths.table("features")
    with path.open("ab") as output:
        output.write(b"changed")
    with pytest.raises(ValueError, match="Modified"):
        daily_data.read_stage(context, "features")
    with pytest.raises(ValueError, match="Stale"):
        daily_data.read_stage(context, "dataset")


def test_nullable_and_empty_tables_roundtrip(tmp_path):
    frame = pd.DataFrame(
        {
            "target": pd.Series([1, pd.NA], dtype="Int8"),
            "flag": pd.Series([False, pd.NA], dtype="boolean"),
            "error": ["", "NA"],
            "time": pd.to_datetime(["2020-01-01", None], utc=True),
            "value": [1.2345678901234567, float("nan")],
        }
    )
    for name, table in (("typed", frame), ("empty", pd.DataFrame())):
        path = tmp_path / (name + ".csv.gz")
        write_table(table, path, root=tmp_path, context={}, sources=[])
        actual = read_table(path, context={}, stage="test")
        pd.testing.assert_frame_equal(actual, table, check_exact=True)


def test_workbook_values_and_export_does_not_recalculate(context, monkeypatch):
    daily_data.prepare(context)
    workbook = load_workbook(
        context.paths.dataset_workbook, read_only=True, data_only=True
    )
    assert workbook.sheetnames == [
        "Summary",
        "Aligned data",
        "Features",
        "Targets",
        "Modeling dataset",
        "Definitions",
    ]
    expected = daily_data.read_stage(context, "dataset")
    sheet = workbook["Modeling dataset"]
    assert sheet.max_row == len(expected) + 1
    rows = list(sheet.values)
    assert rows[0] == (expected.index.name, *expected.columns)
    assert rows[1][0] == expected.index[0].isoformat()
    for actual, value in zip(rows[1][1:], expected.iloc[0], strict=True):
        if isinstance(value, pd.Timestamp):
            assert actual == value.isoformat()
        elif isinstance(value, (int, float)) and not isinstance(value, bool):
            assert actual == pytest.approx(value)
        else:
            assert actual == value
    workbook.close()
    monkeypatch.setattr(
        daily_data,
        "calculate_daily_features",
        lambda *args: pytest.fail("Unexpected feature calculation"),
    )
    monkeypatch.setattr(
        daily_data,
        "build_next_day_d0_targets",
        lambda *args: pytest.fail("Unexpected target calculation"),
    )
    daily_data.export_dataset(context, overwrite=True)


def test_pilot_results_can_be_exported_without_training(context, monkeypatch):
    daily_data.prepare(context)
    paths = svm.run(context)
    assert set(paths) == {"predictions", "summary"}
    workbook_path = context.paths.results("pilot") / "svm_results.xlsx"
    workbook = load_workbook(workbook_path, read_only=True)
    assert "Rolling predictions" in workbook.sheetnames
    workbook.close()
    monkeypatch.setattr(
        svm,
        "run_default_svc_rolling_pilot",
        lambda *args, **kwargs: pytest.fail("Unexpected model fit"),
    )
    svm.export_results(context, overwrite=True)
    with pytest.raises(FileExistsError):
        svm.run(context)


def test_small_grid_pipeline_matches_in_memory(context):
    from test_daily_svm_grid_search import _small_config

    from tcc_crypto.data.development import build_symbol_development_dataset
    from tcc_crypto.modeling.grid_search import run_temporal_grid_search

    config = replace(
        _small_config(),
        raw_directory=context.paths.raw.parent,
        development_end_utc=context.data.sample_end_utc,
    )
    context = replace(
        context,
        profile="development",
        model=config,
        data=replace(context.data, require_warmup_start=False),
    )
    expected = build_symbol_development_dataset(config, context.data.symbol).dataset
    daily_data.prepare(context)
    actual = daily_data.read_stage(context, "dataset")
    pd.testing.assert_frame_equal(actual, expected, check_freq=False, check_exact=True)
    reference = run_temporal_grid_search(
        expected,
        symbol=context.data.symbol,
        config=config,
        selection_config=context.features.feature_selection,
    )
    paths = svm.run(context)
    for name, path in paths.items():
        saved = read_table(path, context=context.signature(), stage="grid-search")
        compared = getattr(reference, name)
        pd.testing.assert_frame_equal(
            saved.drop(columns=["fit_seconds", "mean_fit_seconds"], errors="ignore"),
            compared.drop(columns=["fit_seconds", "mean_fit_seconds"], errors="ignore"),
            check_exact=True,
        )
    book = load_workbook(
        context.paths.results("grid-search") / "svm_results.xlsx", read_only=True
    )
    assert "Candidate ranking" in book.sheetnames
    book.close()


def test_cli_paths_outside_repository_without_creating_outputs(tmp_path):
    environment = dict(__import__("os").environ, PYTHONPATH=str(ROOT / "src"))
    result = subprocess.run(
        [sys.executable, "-m", "tcc_crypto.daily", "paths", "--profile", "pilot"],
        cwd=tmp_path,
        env=environment,
        capture_output=True,
        text=True,
        check=True,
    )
    paths = json.loads(result.stdout)
    assert paths["features"].startswith(
        str(ROOT / "data/processed/binance/spot/1d/features")
    )
    assert paths["results_workbook"].endswith("svm_results.xlsx")
    assert list(tmp_path.iterdir()) == []


def test_legacy_explicit_output_directory_is_respected(tmp_path):
    path = tmp_path / "legacy.json"
    payload = json.loads(
        (ROOT / "configs/binance_daily_svm_pilot_dataset.json").read_text()
    )
    payload["output_directory"] = "outputs/custom_pilot"
    path.write_text(json.dumps(payload))
    context = load_context("pilot", config_path=path, root=ROOT)
    assert context.paths.results("pilot") == ROOT / "outputs/custom_pilot/BTCUSDT"


def test_ineligible_grid_produces_auditable_empty_results(context):
    from test_daily_svm_grid_search import _small_config

    config = replace(_small_config(), excluded_symbols=(context.data.symbol,))
    context = replace(
        context,
        profile="development",
        model=config,
        data=replace(context.data, require_warmup_start=False),
    )
    daily_data.prepare(context)
    paths = svm.run(context)
    for name, path in paths.items():
        table = read_table(path, context=context.signature(), stage="grid-search")
        if name == "eligibility":
            assert set(table.status) == {"ineligible"}
        else:
            assert table.empty


def test_serialized_context_rejects_changed_raw_source(context):
    daily_data.prepare(context)
    raw = daily_data.raw_source(context)
    with raw.open("ab") as output:
        output.write(b"modified")
    with pytest.raises(ValueError, match="Stale"):
        daily_data.read_stage(context, "dataset")


def test_metadata_roundtrip_preserves_mixed_candidate_parameters(tmp_path):
    table = pd.DataFrame(
        {
            "gamma": pd.Series([pd.NA, "scale", 0.1], dtype=object),
            "degree": pd.Series([pd.NA, pd.NA, 3], dtype=object),
            "class_weight": pd.Series([None, "balanced", None], dtype=object),
        }
    )
    path = tmp_path / "parameters.csv.gz"
    write_table(table, path, root=tmp_path, context={}, sources=[])
    actual = read_table(path, context={}, stage="test")
    pd.testing.assert_frame_equal(actual, table, check_exact=True)


def test_notebooks_compile_and_preserve_kernel_metadata():
    for name in (
        "13_binance_daily_feature_selection_analysis",
        "14_binance_daily_svm_pilot_walkthrough",
        "15_binance_daily_svm_grid_search_walkthrough",
    ):
        notebook = json.loads((ROOT / "notebooks" / (name + ".ipynb")).read_text())
        for position, cell in enumerate(notebook["cells"]):
            if cell["cell_type"] == "code":
                compile("".join(cell["source"]), f"{name}:{position}", "exec")
        assert notebook["metadata"]["kernelspec"]["name"] == "python3"


def test_each_cli_preparation_stage_runs_from_another_directory(context, tmp_path):
    import os

    feature_path = tmp_path / "features.json"
    feature_payload = json.loads(context.configuration_files[1].read_text())
    feature_payload["collection_config"] = str(context.configuration_files[2])
    feature_path.write_text(json.dumps(feature_payload))
    pilot_path = tmp_path / "pilot.json"
    pilot_payload = json.loads(context.configuration_files[0].read_text())
    pilot_payload["feature_config"] = str(feature_path)
    pilot_payload["raw_directory"] = str(context.paths.raw.parent)
    pilot_path.write_text(json.dumps(pilot_payload))
    environment = dict(os.environ, PYTHONPATH=str(ROOT / "src"))
    for stage in ("align", "features", "targets", "dataset"):
        subprocess.run(
            [
                sys.executable,
                "-m",
                "tcc_crypto.daily",
                stage,
                "--profile",
                "pilot",
                "--config",
                str(pilot_path),
                "--repository-root",
                str(tmp_path),
            ],
            cwd=tmp_path,
            env=environment,
            capture_output=True,
            text=True,
            check=True,
        )
    resolved = load_context("pilot", config_path=pilot_path, root=tmp_path)
    data = daily_data.read_stage(resolved, "dataset")
    assert data.iloc[-1].target_2pct is pd.NA
    assert resolved.paths.dataset_workbook.exists()


def test_eligibility_only_persists_without_model_fit(context, monkeypatch):
    from test_daily_svm_grid_search import _small_config

    context = replace(context, profile="development", model=_small_config())
    daily_data.prepare(context)
    monkeypatch.setattr(
        svm,
        "run_temporal_grid_search",
        lambda *args, **kwargs: pytest.fail("Unexpected fit"),
    )
    path = svm.assess(context)
    assert path.exists()
    assert not (path.parent / "fold_results.csv.gz").exists()
