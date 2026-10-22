"""Behavioral references and independent-stage tests for exploratory research."""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import sys
from dataclasses import replace
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from tcc_crypto.analysis import daily_descriptive as daily
from tcc_crypto.analysis import hourly_activity, hourly_prices
from tcc_crypto.analysis.correlation import (
    calculate_correlations,
    load_correlation_config,
)
from tcc_crypto.analysis.market_exploration import (
    annual_volume,
    monthly_table,
    trade_counts,
)
from tcc_crypto.collection.exploratory import discover_symbols, exploratory_klines
from tcc_crypto.reporting.migration import MigrationItem, apply_migration, inventory
from tcc_crypto.reporting.provenance import file_sha256
from tcc_crypto.targets.future_high import calculate_future_high_returns
from tcc_crypto.workflows import exploration as workflow
from tcc_crypto.workflows.exploration_context import API_ANALYSES, DEFAULT_CONFIGS

ROOT = Path(__file__).resolve().parents[1]
FIXTURES = ROOT / "tests/fixtures/exploration_refactor"


def golden(name):
    return pd.read_json(FIXTURES / (name + ".json"), orient="table")


def test_old_statistical_and_outcome_references():
    values = golden("input")
    x = np.arange(len(values), dtype=float) + 2
    activity = pd.DataFrame(
        {k: x * (i + 1) for i, k in enumerate(hourly_activity.ACTIVITY_FIELDS)},
        index=values.index,
    )
    outputs = {
        "hourly_prices": hourly_prices.describe_price_levels(
            "TESTUSDT", values[["open", "high", "low", "close"]]
        ),
        "hourly_activity": hourly_activity.describe_market_activity(
            "TESTUSDT", activity
        ),
        "daily_statistics": daily.describe_statistics(values),
        "daily_semesters": daily.describe_semesters(
            values,
            analysis_start_utc=values.index[0],
            analysis_end_utc=values.index[-1],
        ),
        "daily_extrema": daily.describe_extrema("TESTUSDT", values),
        "outcomes": calculate_future_high_returns(
            values, start_utc=values.index[0], end_utc=values.index[-1]
        ),
    }
    for name, actual in outputs.items():
        pd.testing.assert_frame_equal(
            actual,
            golden(name),
            check_freq=False,
            check_exact=False,
            rtol=1e-12,
            atol=1e-12,
        )


def test_original_monthly_and_daily_aggregation_references():
    raw = json.loads((FIXTURES / "market_raw.json").read_text())
    collection = {
        "execution_time": raw["execution_time"],
        "cutoff": raw["cutoff"],
        "records": [
            {"pair": raw["pair"], "records": {"1M": raw["monthly"], "1d": raw["daily"]}}
        ],
    }
    url = "https://api.binance.com/api/v3/klines"
    monthly = monthly_table(collection, url)
    annual = annual_volume(
        monthly, execution_time=pd.Timestamp(raw["execution_time"]), source_url=url
    )
    counts = trade_counts(
        collection, start_utc=pd.Timestamp("2017-01-01T00:00:00Z"), source_url=url
    )
    for name, frame in [
        ("monthly", monthly),
        ("annual", annual),
        ("trade_counts", counts),
    ]:
        pd.testing.assert_frame_equal(
            frame,
            golden(name),
            check_dtype=False,
            check_exact=False,
            rtol=1e-12,
            atol=1e-12,
        )
    assert monthly.iloc[-1]["Month Type"] == "Partial month"
    assert counts.iloc[0]["Missing Days Within Observed Span"] == 1


def test_monthly_pagination_keeps_partial_and_advances_by_open_time():
    raw = json.loads((FIXTURES / "market_raw.json").read_text())
    calls = []

    def request(endpoint, params):
        calls.append(dict(params))
        return [row for row in raw["monthly"] if row[0] >= params["startTime"]][:1]

    rows = exploratory_klines(
        "TESTBTC",
        "1M",
        start_utc=pd.Timestamp("2025-12-01T00:00:00Z"),
        end_utc=pd.Timestamp(raw["execution_time"]),
        requester=request,
        limit=1,
        closed_only=False,
    )
    assert rows == raw["monthly"]
    assert calls[1]["startTime"] == raw["monthly"][0][0] + 1
    assert (
        len(
            exploratory_klines(
                "TESTBTC",
                "1M",
                start_utc=pd.Timestamp("2025-12-01T00:00:00Z"),
                end_utc=pd.Timestamp(raw["execution_time"]),
                requester=request,
                limit=1,
                closed_only=True,
            )
        )
        == 1
    )
    assert (
        exploratory_klines(
            "TESTBTC",
            "1M",
            start_utc=pd.Timestamp("2025-12-01T00:00:00Z"),
            end_utc=pd.Timestamp(raw["execution_time"]),
            requester=lambda *args: [],
            limit=1000,
            closed_only=False,
        )
        == []
    )
    with pytest.raises(ValueError, match="backwards"):
        exploratory_klines(
            "TESTBTC",
            "1M",
            start_utc=pd.Timestamp("2025-12-01T00:00:00Z"),
            end_utc=pd.Timestamp(raw["execution_time"]),
            requester=lambda *args: [raw["monthly"][0]],
            limit=1,
            closed_only=False,
        )


def test_catalog_preserves_status_and_spot_filter_difference():
    symbols = [
        {
            "symbol": "ABTC",
            "baseAsset": "A",
            "quoteAsset": "BTC",
            "status": "BREAK",
            "isSpotTradingAllowed": True,
        },
        {
            "symbol": "BBTC",
            "baseAsset": "B",
            "quoteAsset": "BTC",
            "status": "BREAK",
            "isSpotTradingAllowed": False,
        },
    ]
    assert (
        len(discover_symbols({"symbols": symbols}, quote_asset="BTC", spot_only=True))
        == 1
    )
    assert (
        len(discover_symbols({"symbols": symbols}, quote_asset="BTC", spot_only=False))
        == 2
    )


def test_correlations_use_exact_pairwise_observations_and_frozen_universe():
    dates = pd.date_range("2020-01-01", periods=5, freq="h", tz="UTC")
    series = pd.DataFrame(
        {
            "BTCUSDT": [1.0, 2.0, 3.0, 4.0, 5.0],
            "OTHER": [2.0, np.nan, 4.0, 8.0, np.nan],
        },
        index=dates,
    )
    result = calculate_correlations(series, "BTCUSDT")
    assert result["paired_observations"].loc["BTCUSDT", "OTHER"] == 3
    assert result["correlations"].loc["BTCUSDT", "OTHER"] == pytest.approx(
        series.dropna().corr().iloc[0, 1]
    )
    for name in ("open", "high_open_return"):
        cfg = load_correlation_config(
            ROOT / "configs" / f"binance_hourly_{name}_correlation.json"
        )
        assert len(cfg.candidate_symbols) == 30
        assert "HBARUSDT" not in cfg.candidate_symbols


@pytest.fixture
def research_root(tmp_path):
    shutil.copytree(ROOT / "configs", tmp_path / "configs")
    (tmp_path / "pyproject.toml").write_text('[project]\nname="test"\n')
    for interval in ("1h", "1d"):
        index = pd.date_range(
            "2019-07-01", periods=130, freq="h" if interval == "1h" else "D", tz="UTC"
        )
        for symbol in ("BTCUSDT", "ETHUSDT", "LUNAUSDT"):
            x = np.arange(len(index))
            base = 100 + np.sin(x) * 4 + x / 10
            values = pd.DataFrame(
                {
                    "open_time_ms": index.asi8 // 1_000_000,
                    "open": base,
                    "high": base + 1 + (x % 6),
                    "low": base - 1,
                    "close": base + 0.5,
                    "quote_asset_volume": (x + 1) * 100.0,
                    "base_asset_volume": x + 1.0,
                    "number_of_trades": x + 1,
                    "taker_buy_base_asset_volume": x * 0.2,
                    "taker_buy_quote_asset_volume": x * 10.0,
                }
            )
            if symbol == "LUNAUSDT":
                values = values.drop(30)
            path = (
                tmp_path
                / f"data/raw/binance/spot/{interval}"
                / symbol
                / f"{symbol}_{interval}_sample.csv.gz"
            )
            path.parent.mkdir(parents=True, exist_ok=True)
            values.to_csv(path, index=False)
    return tmp_path


@pytest.mark.parametrize(
    "analysis",
    [
        "hourly-prices",
        "hourly-activity",
        "daily-descriptive",
        "future-high",
        "feature-selection",
    ],
)
def test_staged_roundtrip_and_export_without_recalculation(
    research_root, analysis, monkeypatch
):
    symbols = (
        ("BTCUSDT", "ETHUSDT", "LUNAUSDT")
        if analysis == "feature-selection"
        else ("BTCUSDT", "ETHUSDT")
    )
    context = workflow.open_exploration(
        analysis, root=research_root, run_id="reference", symbols=symbols
    )
    if analysis == "future-high":
        context = replace(
            context, settings=replace(context.settings, additional_symbol_count=1)
        )
    raw_before = {
        path: file_sha256(path) for path in (research_root / "data/raw").rglob("*.gz")
    }
    inputs = workflow.load_inputs(context)
    workflow.save_inputs(context, inputs)
    actual = workflow.calculate(context, inputs)
    workflow.save_results(context, actual)
    restored = workflow.load_results(context)
    for name, frame in actual.tables.items():
        pd.testing.assert_frame_equal(frame, restored.tables[name], check_exact=True)
    monkeypatch.setattr(
        workflow, "calculate", lambda *a, **k: pytest.fail("Unexpected calculation")
    )
    workbooks = workflow.export_workbooks(
        context, workflow.load_saved_inputs(context), restored
    )
    for path in workbooks.values():
        book = load_workbook(path)
        if analysis == "hourly-prices":
            assert book.sheetnames == [
                "Price Statistics",
                "Coverage",
                "Extrema",
                "Definitions",
            ]
        if analysis == "hourly-activity":
            assert book.sheetnames == [
                "Activity Statistics",
                "Coverage",
                "Extrema",
                "Definitions",
            ]
        if analysis == "daily-descriptive":
            assert len(book.sheetnames) == 5
            assert len(book["Overview"]._charts) == 2
            assert book["Chart Data"]["A2"].data_type == "d"
            assert (
                book["Overview"]["A2"].value == actual.tables["statistics"].iloc[0, 0]
            )
        book.close()
    if analysis == "feature-selection":
        assert (
            restored.tables["manifest"].set_index("symbol").loc["LUNAUSDT", "status"]
            == "BLOCKED"
        )
        book = load_workbook(workbooks["LUNAUSDT"])
        assert "Missing dates" in book.sheetnames
        book.close()
    with pytest.raises(FileExistsError):
        workflow.export_workbooks(context, inputs, restored)
    with pytest.raises(FileExistsError):
        workflow.save_results(context, actual)
    assert raw_before == {path: file_sha256(path) for path in raw_before}
    source = inputs.sources[0]
    source.write_bytes(source.read_bytes() + b"changed")
    with pytest.raises(ValueError, match="Stale|modified"):
        workflow.load_results(context)


def test_missing_and_legacy_configuration_fail_clearly(research_root):
    context = workflow.open_exploration(
        "hourly-prices", root=research_root, run_id="absent", symbols=("BTCUSDT",)
    )
    with pytest.raises(FileNotFoundError, match="Missing"):
        workflow.load_results(context)
    p = research_root / "configs/binance_hourly_price_descriptive_analysis.json"
    payload = json.loads(p.read_text())
    payload["output_directory"] = "outputs/old"
    p.write_text(json.dumps(payload))
    with pytest.raises(ValueError, match="Legacy analysis path"):
        workflow.open_exploration("hourly-prices", root=research_root)


def fake_request(endpoint, params):
    if endpoint.endswith("exchangeInfo"):
        return {
            "symbols": [
                {
                    "symbol": "TESTBTC",
                    "baseAsset": "TEST",
                    "quoteAsset": "BTC",
                    "status": "BREAK",
                    "isSpotTradingAllowed": True,
                },
                {
                    "symbol": "TESTUSDT",
                    "baseAsset": "TEST",
                    "quoteAsset": "USDT",
                    "status": "TRADING",
                    "isSpotTradingAllowed": True,
                },
            ]
        }
    raw = json.loads((FIXTURES / "market_raw.json").read_text())
    if params["limit"] == 1:
        return [raw["daily"][0]]
    return raw["monthly"] if params["interval"] == "1M" else raw["daily"]


@pytest.mark.parametrize("analysis", API_ANALYSES)
def test_api_collection_reuses_snapshots_and_roundtrips(research_root, analysis):
    context = workflow.open_exploration(analysis, root=research_root, run_id="api")
    execution = pd.Timestamp("2026-01-15T12:00:00Z")
    inputs = workflow.collect_inputs(
        context, requester=fake_request, execution_time=execution
    )
    repeated = workflow.collect_inputs(
        context,
        requester=lambda *a: pytest.fail("Snapshot should prevent network access"),
    )
    assert repeated.raw == inputs.raw
    with pytest.raises(ValueError, match="cutoff"):
        workflow.collect_inputs(
            context,
            requester=fake_request,
            execution_time=execution + pd.Timedelta(days=1),
        )
    workflow.save_inputs(context, inputs)
    result = workflow.calculate(context, inputs)
    workflow.save_results(context, result)
    for name, frame in result.tables.items():
        pd.testing.assert_frame_equal(
            frame, workflow.load_results(context).tables[name]
        )
    assert workflow.export_workbooks(
        context, workflow.load_saved_inputs(context), workflow.load_results(context)
    )


def test_cli_paths_and_independent_calculate_export_from_other_directory(research_root):
    env = dict(
        os.environ, PYTHONPATH=str(ROOT / "src"), MPLCONFIGDIR="/tmp/tcc_remaining_mpl"
    )
    command = [sys.executable, "-m", "tcc_crypto.explore", "hourly-prices"]
    args = [
        "--repository-root",
        str(research_root),
        "--run-id",
        "cli",
        "--symbols",
        "BTCUSDT",
    ]
    subprocess.run(
        [*command, "paths", *args], cwd="/tmp", env=env, capture_output=True, check=True
    )
    assert not (research_root / "outputs").exists()
    for operation in ("calculate", "export"):
        subprocess.run(
            [*command, operation, *args],
            cwd="/tmp",
            env=env,
            capture_output=True,
            check=True,
        )
    context = workflow.open_exploration(
        "hourly-prices", root=research_root, run_id="cli", symbols=("BTCUSDT",)
    )
    assert workflow.load_results(context).tables["statistics"].shape[0] == 4


def test_migration_dry_run_collision_resume_and_personal_exclusion(research_root):
    source = (
        research_root
        / "outputs/binance_hourly_price_descriptive_analysis/statistics.csv"
    )
    source.parent.mkdir(parents=True)
    source.write_text("value\n1\n")
    personal = research_root / "outputs/pdf/certificate.pdf"
    personal.parent.mkdir(parents=True)
    personal.write_bytes(b"personal")
    items = inventory(research_root)
    assert len(items) == 1
    assert source.exists()
    assert not (research_root / "outputs/migrations").exists()
    destination = research_root / items[0].destination
    destination.parent.mkdir(parents=True)
    destination.write_bytes(b"collision")
    with pytest.raises(FileExistsError):
        apply_migration(research_root, items)
    assert source.exists()
    destination.write_bytes(source.read_bytes())
    result = apply_migration(research_root, items)
    assert all(i["status"] == "migrated" for i in result["items"])
    assert not source.exists()
    assert personal.read_bytes() == b"personal"
    assert apply_migration(research_root) == result
    assert file_sha256(destination) == items[0].sha256


@pytest.mark.parametrize("number", range(1, 16))
def test_each_notebook_executes_with_local_inputs_and_simulated_network(
    research_root, monkeypatch, number
):
    """Execute the actual code cells, reducing only smoke-test data and model scopes."""
    import matplotlib.pyplot as plt
    from IPython import display as ipython_display
    from test_daily_pipeline import context as existing_daily_fixture
    from test_daily_svm_grid_search import _small_config

    from tcc_crypto.collection import client
    from tcc_crypto.collection import exploratory as exploratory_collection
    from tcc_crypto.workflows import collection_review, daily_review

    real_open = workflow.open_exploration

    def open_local(analysis, **kwargs):
        symbols = (
            ("BTCUSDT", "ETHUSDT", "LUNAUSDT")
            if analysis == "feature-selection"
            else ("BTCUSDT", "ETHUSDT")
        )
        if analysis in API_ANALYSES or analysis in (
            "open-correlation",
            "return-correlation",
        ):
            symbols = None
        if analysis in ("open-correlation", "return-correlation"):
            cfg = load_correlation_config(
                research_root / "configs" / DEFAULT_CONFIGS[analysis]
            )
            source = (
                research_root
                / "data/raw/binance/spot/1h/BTCUSDT/BTCUSDT_1h_sample.csv.gz"
            )
            for symbol in cfg.candidate_symbols:
                path = (
                    research_root
                    / f"data/raw/binance/spot/1h/{symbol}/{symbol}_1h_sample.csv.gz"
                )
                if not path.exists():
                    path.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copyfile(source, path)
        context = real_open(
            analysis, root=research_root, run_id="notebook", symbols=symbols
        )
        if analysis == "future-high":
            context = replace(
                context, settings=replace(context.settings, additional_symbol_count=1)
            )
        return context

    monkeypatch.setattr(workflow, "open_exploration", open_local)
    monkeypatch.setattr(
        client, "request_json", lambda url, params, **kwargs: fake_request(url, params)
    )
    monkeypatch.setattr(
        exploratory_collection,
        "request_json",
        lambda url, params, **kwargs: fake_request(url, params),
    )
    # Existing collector uses the same client; exercise real persistence with two valid candles.
    if number in (5, 10):
        monkeypatch.setattr(
            collection_review, "repository_root", lambda root=None: research_root
        )
        original = collection_review.configuration

        def collection_config(interval, **kwargs):
            config = original(interval, root=research_root)
            start = pd.Timestamp("2026-01-01T00:00:00Z")
            end = (
                start + pd.Timedelta(days=2) - pd.Timedelta(milliseconds=1)
                if interval == "1d"
                else start + pd.Timedelta(hours=2) - pd.Timedelta(milliseconds=1)
            )
            return replace(
                config, symbols=("SMOKEUSDT",), start_time_utc=start, end_time_utc=end
            )

        monkeypatch.setattr(collection_review, "configuration", collection_config)

        def collection_response(url, params, **kwargs):
            start = pd.Timestamp("2026-01-01T00:00:00Z")
            delta = (
                pd.Timedelta(days=1)
                if params["interval"] == "1d"
                else pd.Timedelta(hours=1)
            )
            rows = []
            for i in range(2):
                date = start + i * delta
                end = date + delta - pd.Timedelta(milliseconds=1)
                rows.append(
                    [
                        date.value // 1000000,
                        "10",
                        "12",
                        "9",
                        "11",
                        "1",
                        end.value // 1000000,
                        "10",
                        3,
                        "0",
                        "0",
                        "0",
                    ]
                )
            return rows

        monkeypatch.setattr(client, "request_json", collection_response)
    if number in (14, 15):
        context = existing_daily_fixture.__wrapped__(research_root)
        if number == 15:
            context = replace(context, profile="development", model=_small_config())
        monkeypatch.setattr(
            daily_review, "load_context", lambda *args, **kwargs: context
        )
    monkeypatch.setattr(ipython_display, "display", lambda *args, **kwargs: None)
    path = next((ROOT / "notebooks").glob(f"{number:02d}_*.ipynb"))
    namespace = {"__name__": "__notebook__"}
    try:
        for index, cell in enumerate(json.loads(path.read_text())["cells"]):
            if cell["cell_type"] == "code":
                exec(
                    compile(
                        "".join(cell["source"]), f"{path.name}:cell{index}", "exec"
                    ),
                    namespace,
                )
    finally:
        plt.close("all")


def test_missing_feature_source_is_blocked_and_exportable(research_root):
    context = workflow.open_exploration(
        "feature-selection",
        root=research_root,
        run_id="missing-symbol",
        symbols=("FLOKIUSDT",),
    )
    inputs = workflow.load_inputs(context)
    assert inputs.values == {}
    assert inputs.diagnostics.iloc[0]["status"] == "BLOCKED"
    workflow.save_inputs(context, inputs)
    result = workflow.calculate(context, inputs)
    workflow.save_results(context, result)
    books = workflow.export_workbooks(context, inputs, result)
    book = load_workbook(books["FLOKIUSDT"])
    assert "Missing dates" in book.sheetnames
    book.close()


def test_request_validation_and_migration_reject_changed_or_raw_sources(research_root):
    from tcc_crypto.collection.exploratory import validate_exploratory_rows

    raw = json.loads((FIXTURES / "market_raw.json").read_text())
    row = list(raw["monthly"][0])
    row[2] = "8"
    with pytest.raises(ValueError, match="OHLC"):
        validate_exploratory_rows([row], "1M")
    source = research_root / "data/raw/protected.csv"
    source.write_text("protected")
    item = MigrationItem(
        str(source.relative_to(research_root)),
        "outputs/history/protected.csv",
        file_sha256(source),
    )
    with pytest.raises(ValueError, match="Unsafe"):
        apply_migration(research_root, [item])
    assert source.read_text() == "protected"


def test_migration_adds_new_recognized_files_to_the_existing_journal(research_root):
    directory = research_root / "outputs/binance_hourly_price_descriptive_analysis"
    directory.mkdir(parents=True)
    first = directory / "first.csv"
    first.write_text("first")
    assert len(apply_migration(research_root)["items"]) == 1
    directory.mkdir(parents=True, exist_ok=True)
    second = directory / "second.csv"
    second.write_text("second")
    assert len(apply_migration(research_root)["items"]) == 2
    assert not second.exists()
